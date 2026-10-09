"""Учёт численности: загрузка кодификатора видов работ и выгрузки факта из SharePoint (2026-10-09, Docs/headcount.md §5).

Два файла заказчика:

1. «Справочник по видам работ (PBI).xlsx», лист `РаботыНаименовКодифМСУ` — общий кодификатор. Ключ — код. Загрузка добавляет и
   обновляет записи, но НЕ выводит из действия коды, которых в файле нет (файл мог быть неполным): их число только сообщается.
2. «Выгрузка факт численности 2026.xlsx» — строки «дата ввода / число / день / код / ИНН / GUID объекта 1С / название объекта 1С».

Выгрузка идёт в два шага. `analyze` ничего не пишет: показывает объекты 1С (GUID + название), что из них уже сопоставлено с объектами
системы (по GUID или точному названию), неизвестные коды и грязные данные. Объекты сопоставляет человек вручную при первой загрузке
(решение 2026-10-09); `apply` принимает сопоставление {GUID: id объекта | null — пропустить}, запоминает его в `objects.guid_1c` и
грузит строки. Файл при `apply` читается заново, но это тот же файл: сопоставление привязано к GUID, а не к строкам.

Правила загрузки строк:
- Ключ — (объект, день, подрядчик, вид работ). Строки с одним ключом сворачиваются: актуальна последняя по времени ввода, все значения
  попадают в историю в хронологическом порядке (SharePoint допускал повторы, у 41 из 176 повторов числа различались).
- Подрядчик — по ИНН внутри объекта. Значение в колонке ИНН, не состоящее из цифр (бывает название организации), становится
  названием подрядчика без ИНН. ИНН других длин грузятся как есть и в списках получают пометку «ИНН не проверен».
- Запись, у которой в системе уже есть правка из формы (`source = 'form'`), НЕ перезаписывается: данные, внесённые людьми, главнее
  выгрузки. Повторная загрузка того же файла ничего не меняет (идемпотентна).
- Просрочка (`late`) у загруженных записей не ставится: срок сверяется только для внесения в системе.
"""

import io
import json
import re
import sqlite3
from collections import defaultdict
from datetime import date, datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook

from app import activity
from app.access import require_service_feature
from app.db import begin_write, get_connection
from app.headcount import MAX_WORKERS, MSK, clean_inn, find_counterparty_by_inn, inn_status
from app.upload_limits import read_upload_limited

router = APIRouter(prefix="/headcount", tags=["headcount"])

CODIFIER_SHEET = "РаботыНаименовКодифМСУ"
GUID_RE = re.compile(r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$")

CODIFIER_COLUMNS = {
    "code": "Кодификатор", "name": "Наименование работ по кодификатору", "parent": "Родитель по кодификатору",
    "section": "Раздел по кодификатору", "unit": "Ед. измерения",
}
HISTORY_COLUMNS = {
    "entered": "mcyDateOrig", "workers": "mcyNumberWorkers", "day": "Дата", "code": "mcyCodifier", "inn": "inn",
    "guid": "guid1c", "name_1c": "НаименованиеОбъекта1С",
}


# ------------------------------------------------------------------ чтение файлов

def _norm(value) -> str:
    return " ".join(str(value or "").split()).lower()


def _sheet_rows(file_bytes: bytes, prefer: Optional[str] = None) -> list:
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Файл не читается как Excel (.xlsx)")
    try:
        ws = wb[prefer] if prefer and prefer in wb.sheetnames else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _columns(header: list, wanted: dict, file_label: str) -> dict:
    index = {_norm(v): i for i, v in enumerate(header) if v is not None}
    missing = [title for title in wanted.values() if _norm(title) not in index]
    if missing:
        raise ValueError(f"{file_label}: нет колонок {', '.join('«%s»' % m for m in missing)}")
    return {key: index[_norm(title)] for key, title in wanted.items()}


def _text(value) -> Optional[str]:
    value = " ".join(str(value if value is not None else "").split())
    return value or None


def _as_date(value) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except (TypeError, ValueError):
        return None


def _as_datetime(value) -> Optional[datetime]:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    try:
        return datetime.fromisoformat(str(value))
    except (TypeError, ValueError):
        return None


def _inn_cell(value) -> Optional[str]:
    """Ячейка ИНН как текст: Excel мог отдать число (7745000111, 7745000111.0)."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return _text(value)


# ------------------------------------------------------------------ кодификатор

def import_codifier(conn: sqlite3.Connection, file_bytes: bytes, dry_run: bool = False) -> dict:
    rows = _sheet_rows(file_bytes, CODIFIER_SHEET)
    if not rows:
        raise ValueError("Файл пуст")
    col = _columns(rows[0], CODIFIER_COLUMNS, "Кодификатор")
    items, seen = [], set()
    for order, row in enumerate(rows[1:], start=1):
        code = _text(row[col["code"]] if col["code"] < len(row) else None)
        name = _text(row[col["name"]] if col["name"] < len(row) else None)
        if not code or not name:
            continue
        if code in seen:
            raise ValueError(f"Код «{code}» встречается в файле дважды — ключ кодификатора должен быть уникален")
        seen.add(code)
        pick = lambda key: _text(row[col[key]]) if col[key] < len(row) else None
        items.append((code, name, pick("parent"), pick("section"), pick("unit"), order))
    existing = {r["code"]: r for r in conn.execute("SELECT * FROM work_codifier")}
    created = updated = unchanged = 0
    for code, name, parent, section, unit, order in items:
        cur = existing.get(code)
        if cur is None:
            created += 1
        elif (cur["name"], cur["parent_name"], cur["section_name"], cur["unit"], cur["sort_order"], cur["retired_at"]) == \
                (name, parent, section, unit, order, None):
            unchanged += 1
        else:
            updated += 1
    absent = sum(1 for code, r in existing.items() if code not in seen and not r["retired_at"])
    if not dry_run:
        begin_write(conn)
        for code, name, parent, section, unit, order in items:
            conn.execute(
                "INSERT INTO work_codifier (code, name, parent_name, section_name, unit, sort_order) VALUES (?, ?, ?, ?, ?, ?) "
                "ON CONFLICT(code) DO UPDATE SET name = excluded.name, parent_name = excluded.parent_name, "
                "section_name = excluded.section_name, unit = excluded.unit, sort_order = excluded.sort_order, "
                "retired_at = NULL, updated_at = datetime('now') "
                "WHERE name IS NOT excluded.name OR parent_name IS NOT excluded.parent_name OR "
                "section_name IS NOT excluded.section_name OR unit IS NOT excluded.unit OR "
                "sort_order IS NOT excluded.sort_order OR retired_at IS NOT NULL",
                (code, name, parent, section, unit, order))
        conn.commit()
    return {"rows": len(items), "created": created, "updated": updated, "unchanged": unchanged,
            "absent_in_file": absent, "dry_run": dry_run}


# ------------------------------------------------------------------ выгрузка факта: разбор

class ParsedHistory:
    """Выгрузка, разобранная в память: валидные строки + счётчики ненужного."""

    def __init__(self):
        self.rows = []              # (entered: datetime, day: date, workers: int, code, inn_cell, guid, name_1c, order)
        self.bad = defaultdict(int)  # причина → число строк
        self.total = 0


def parse_history(file_bytes: bytes) -> ParsedHistory:
    rows = _sheet_rows(file_bytes)
    if not rows:
        raise ValueError("Файл пуст")
    col = _columns(rows[0], HISTORY_COLUMNS, "Выгрузка численности")
    out = ParsedHistory()
    width = max(col.values()) + 1
    for order, row in enumerate(rows[1:]):
        if not any(v not in (None, "") for v in row):
            continue
        out.total += 1
        row = list(row) + [None] * (width - len(row))
        workers = row[col["workers"]]
        if isinstance(workers, float) and workers.is_integer():
            workers = int(workers)
        day = _as_date(row[col["day"]])
        entered = _as_datetime(row[col["entered"]]) or (datetime.combine(day, datetime.min.time()) if day else None)
        code, guid = _text(row[col["code"]]), _text(row[col["guid"]])
        if day is None:
            out.bad["нет даты"] += 1
        elif not isinstance(workers, int) or isinstance(workers, bool) or not 1 <= workers <= MAX_WORKERS:
            out.bad["число рабочих вне 1–100 000"] += 1
        elif not code:
            out.bad["нет кода вида работ"] += 1
        elif not guid or not GUID_RE.match(guid):
            out.bad["нет GUID объекта 1С или он неверного вида"] += 1
        else:
            out.rows.append((entered, day, workers, code, _inn_cell(row[col["inn"]]), guid.lower(),
                             _text(row[col["name_1c"]]), order))
    return out


def _contractor_identity(inn_cell: Optional[str]) -> tuple:
    """(inn_raw, name_raw) подрядчика по ячейке «ИНН». Нецифровое значение — это название организации, а не ИНН."""
    if inn_cell is None:
        return None, None
    inn = clean_inn(inn_cell)
    if inn and inn.isdigit():
        return inn, None
    return None, _text(inn_cell)


# ------------------------------------------------------------------ analyze

def analyze_history(conn: sqlite3.Connection, file_bytes: bytes) -> dict:
    parsed = parse_history(file_bytes)
    codes = {r["code"] for r in conn.execute("SELECT code FROM work_codifier")}
    by_guid = {}
    unknown_codes = defaultdict(int)
    contractors, keys, key_values = set(), defaultdict(list), {}
    for entered, day, workers, code, inn_cell, guid, name_1c, order in parsed.rows:
        g = by_guid.setdefault(guid, {"guid": guid, "name_1c": name_1c, "rows": 0, "date_from": day, "date_to": day, "contractors": set()})
        g["rows"] += 1
        g["date_from"], g["date_to"] = min(g["date_from"], day), max(g["date_to"], day)
        inn, name = _contractor_identity(inn_cell)
        g["contractors"].add(inn or name)
        contractors.add((guid, inn or name))
        if code not in codes:
            unknown_codes[code] += 1
        keys[(guid, day, inn or name, code)].append(workers)
    duplicates = sum(1 for v in keys.values() if len(v) > 1)
    conflicting = sum(1 for v in keys.values() if len(set(v)) > 1)

    objects = [dict(r) for r in conn.execute("SELECT id, name, guid_1c FROM objects ORDER BY name")]
    by_guid_obj = {o["guid_1c"].lower(): o for o in objects if o["guid_1c"]}
    by_name = defaultdict(list)
    for o in objects:
        by_name[_norm(o["name"])].append(o)
    groups = []
    for g in sorted(by_guid.values(), key=lambda x: -x["rows"]):
        match, how = by_guid_obj.get(g["guid"]), "guid"
        if match is None:
            cand = by_name.get(_norm(g["name_1c"]), [])
            match, how = (cand[0], "name") if len(cand) == 1 and not cand[0]["guid_1c"] else (None, None)
        groups.append({"guid": g["guid"], "name_1c": g["name_1c"], "rows": g["rows"], "contractors": len(g["contractors"]),
                       "date_from": g["date_from"].isoformat(), "date_to": g["date_to"].isoformat(),
                       "object_id": match["id"] if match else None, "object_name": match["name"] if match else None,
                       "match": how})
    kinds = defaultdict(int)
    for _, c in contractors:
        if c is None:
            kinds["no_inn"] += 1
        elif c.isdigit():
            kinds[inn_status(c)] += 1
        else:
            kinds["name_in_inn_field"] += 1
    days = sorted({r[1] for r in parsed.rows})
    return {
        "rows_total": parsed.total, "rows_valid": len(parsed.rows), "bad_rows": dict(parsed.bad),
        "date_from": days[0].isoformat() if days else None, "date_to": days[-1].isoformat() if days else None,
        "days": len(days), "weekend_rows": sum(1 for r in parsed.rows if r[1].weekday() >= 5),
        "keys": len(keys), "duplicate_keys": duplicates, "conflicting_keys": conflicting,
        "contractors": {"total": len(contractors), **kinds},
        "unknown_codes": [{"code": c, "rows": n} for c, n in sorted(unknown_codes.items())],
        "objects": groups, "candidates": [{"id": o["id"], "name": o["name"], "guid_1c": o["guid_1c"]} for o in objects],
    }


# ------------------------------------------------------------------ apply

def _utc_stamp(moment: Optional[datetime]) -> str:
    """Время ввода из SharePoint (местное, без пояса) → строка БД в UTC, как `datetime('now')`."""
    if moment is None:
        return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=MSK)
    return moment.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def validate_mapping(conn: sqlite3.Connection, mapping: dict) -> dict:
    """{GUID: id объекта | None} → очищенное. Один объект — один GUID (уникальный индекс idx_objects_guid_1c): объект, уже привязанный
    к ДРУГОМУ GUID, и два GUID на один объект в одном запросе отклоняются с понятным текстом."""
    clean, taken = {}, {}
    for guid, object_id in (mapping or {}).items():
        guid = str(guid).strip().lower()
        if not GUID_RE.match(guid):
            raise HTTPException(status_code=400, detail=f"Неверный GUID 1С: {guid[:40]}")
        if object_id is None:
            clean[guid] = None
            continue
        if not isinstance(object_id, int) or isinstance(object_id, bool):
            raise HTTPException(status_code=400, detail="В сопоставлении ожидается идентификатор объекта или null")
        row = conn.execute("SELECT id, name, guid_1c FROM objects WHERE id = ?", (object_id,)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail=f"Объект {object_id} не найден")
        if row["guid_1c"] and row["guid_1c"].lower() != guid:
            raise HTTPException(status_code=409, detail=f"Объект «{row['name']}» уже привязан к другому GUID 1С")
        other = conn.execute("SELECT name FROM objects WHERE guid_1c = ? AND id <> ?", (guid, object_id)).fetchone()
        if other:
            raise HTTPException(status_code=409, detail=f"GUID 1С уже привязан к объекту «{other['name']}»")
        if object_id in taken and taken[object_id] != guid:
            raise HTTPException(status_code=409, detail=f"Объект «{row['name']}» выбран для двух разных GUID 1С — у объекта он один")
        taken[object_id] = guid
        clean[guid] = object_id
    return clean


def apply_history(conn: sqlite3.Connection, file_bytes: bytes, mapping: dict, user=None, dry_run: bool = False) -> dict:
    parsed = parse_history(file_bytes)
    mapping = validate_mapping(conn, mapping)
    codes = {r["code"]: r["id"] for r in conn.execute("SELECT code, id FROM work_codifier")}
    user_id = user["id"] if user is not None else None

    skipped = defaultdict(int)
    chains = defaultdict(list)   # ключ → [(entered, order, workers)]
    for entered, day, workers, code, inn_cell, guid, name_1c, order in parsed.rows:
        object_id = mapping.get(guid)
        if object_id is None:
            skipped["объект 1С не сопоставлен"] += 1
        elif code not in codes:
            skipped["код вида работ не найден в кодификаторе"] += 1
        else:
            inn, name = _contractor_identity(inn_cell)
            chains[(object_id, day, inn, name, code)].append((entered, order, workers))

    stat = {"records_created": 0, "records_updated": 0, "records_unchanged": 0, "kept_form_records": 0,
            "history_rows": 0, "contractors_created": 0, "contractors_linked_to_counterparty": 0}
    begin_write(conn)
    try:
        for guid, object_id in mapping.items():
            if object_id is not None:
                conn.execute("UPDATE objects SET guid_1c = ? WHERE id = ? AND (guid_1c IS NULL OR guid_1c <> ?)", (guid, object_id, guid))
        contractor_ids = {}

        def contractor(object_id, inn, name) -> int:
            key = (object_id, inn, name)
            if key in contractor_ids:
                return contractor_ids[key]
            row = (conn.execute("SELECT id FROM headcount_contractors WHERE object_id = ? AND inn_raw = ?", (object_id, inn)).fetchone()
                   if inn else
                   conn.execute("SELECT id FROM headcount_contractors WHERE object_id = ? AND inn_raw IS NULL AND name_raw = ?",
                                (object_id, name)).fetchone())
            if row is None:
                cp = find_counterparty_by_inn(conn, inn)
                cur = conn.execute("INSERT INTO headcount_contractors (object_id, counterparty_id, name_raw, inn_raw) VALUES (?, ?, ?, ?)",
                                   (object_id, cp, name, inn))
                stat["contractors_created"] += 1
                stat["contractors_linked_to_counterparty"] += 1 if cp else 0
                contractor_ids[key] = cur.lastrowid
            else:
                contractor_ids[key] = row["id"]
            return contractor_ids[key]

        for (object_id, day, inn, name, code), chain in chains.items():
            chain.sort(key=lambda t: (t[0] or datetime.min, t[1]))
            values, stamps = [], []
            for entered, _, workers in chain:       # подряд идущие одинаковые значения — не изменение
                if not values or values[-1] != workers:
                    values.append(workers)
                    stamps.append(_utc_stamp(entered))
            cid, wid, iso = contractor(object_id, inn, name), codes[code], day.isoformat()
            current = conn.execute(
                "SELECT id, workers, source FROM headcount_records WHERE object_id = ? AND work_date = ? AND contractor_id = ? AND codifier_id = ?",
                (object_id, iso, cid, wid)).fetchone()
            if current is None:
                cur = conn.execute(
                    "INSERT INTO headcount_records (object_id, work_date, contractor_id, codifier_id, workers, late, source, "
                    "entered_at, updated_at) VALUES (?, ?, ?, ?, ?, 0, 'import', ?, ?)",
                    (object_id, iso, cid, wid, values[-1], stamps[0], stamps[-1]))
                record_id, previous = cur.lastrowid, None
                stat["records_created"] += 1
            elif current["source"] == "form":
                stat["kept_form_records"] += 1
                continue
            elif current["workers"] == values[-1]:
                stat["records_unchanged"] += 1
                continue
            else:
                record_id, previous = current["id"], current["workers"]
                conn.execute("UPDATE headcount_records SET workers = ?, updated_at = ? WHERE id = ?", (values[-1], stamps[-1], record_id))
                stat["records_updated"] += 1
            # история: новая запись — вся цепочка значений; уже загруженная и изменившаяся — одна строка «было → стало»
            if previous is None:
                steps = [(None if i == 0 else values[i - 1], values[i], stamps[i]) for i in range(len(values))]
            else:
                steps = [(previous, values[-1], stamps[-1])]
            for old, new, stamp in steps:
                conn.execute(
                    "INSERT INTO headcount_history (record_id, object_id, work_date, contractor_id, codifier_id, old_workers, "
                    "new_workers, source, changed_at, changed_by) VALUES (?, ?, ?, ?, ?, ?, ?, 'import', ?, ?)",
                    (record_id, object_id, iso, cid, wid, old, new, stamp, user_id))
                stat["history_rows"] += 1
        if dry_run:
            conn.rollback()
        else:
            conn.commit()
    except Exception:
        conn.rollback()
        raise
    return {"rows_total": parsed.total, "rows_valid": len(parsed.rows), "bad_rows": dict(parsed.bad),
            "skipped": dict(skipped), "keys": len(chains), **stat, "dry_run": dry_run}


# ------------------------------------------------------------------ API

def _guard(fn):
    try:
        return fn()
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@router.post("/codifier/import")
def api_codifier_import(file: UploadFile = File(...), dry_run: bool = Form(False),
                        admin: sqlite3.Row = Depends(require_service_feature("headcount_admin", "write"))):
    payload = read_upload_limited(file.file)
    conn = get_connection()
    try:
        result = _guard(lambda: import_codifier(conn, payload, dry_run=dry_run))
    finally:
        conn.close()
    if not dry_run:
        activity.log("headcount_codifier_load", user=admin, new_value=file.filename,
                     details={k: result[k] for k in ("rows", "created", "updated", "absent_in_file")})
    return result


@router.post("/import/analyze")
def api_history_analyze(file: UploadFile = File(...),
                        admin: sqlite3.Row = Depends(require_service_feature("headcount_admin", "write"))):
    payload = read_upload_limited(file.file)
    conn = get_connection()
    try:
        return _guard(lambda: analyze_history(conn, payload))
    finally:
        conn.close()


@router.post("/import/apply")
def api_history_apply(file: UploadFile = File(...), mapping: str = Form(..., description='JSON {"GUID": id объекта | null}'),
                      dry_run: bool = Form(False),
                      admin: sqlite3.Row = Depends(require_service_feature("headcount_admin", "write"))):
    try:
        mapping_obj = json.loads(mapping)
        if not isinstance(mapping_obj, dict):
            raise ValueError
    except ValueError:
        raise HTTPException(status_code=400, detail="mapping: ожидается JSON-объект {GUID: id объекта | null}")
    payload = read_upload_limited(file.file)
    conn = get_connection()
    try:
        result = _guard(lambda: apply_history(conn, payload, mapping_obj, user=admin, dry_run=dry_run))
    finally:
        conn.close()
    if not dry_run:
        activity.log("headcount_import", user=admin, new_value=file.filename,
                     details={k: result[k] for k in ("rows_valid", "records_created", "records_updated", "history_rows")})
    return result


@router.put("/objects/{object_id}/guid-1c")
def api_set_guid(object_id: int, body: dict, admin: sqlite3.Row = Depends(require_service_feature("headcount_admin", "write"))):
    """Привязать объект к GUID 1С вручную (или снять привязку: {"guid_1c": null})."""
    raw = body.get("guid_1c")
    guid = None if raw in (None, "") else str(raw).strip().lower()
    if guid is not None and not GUID_RE.match(guid):
        raise HTTPException(status_code=400, detail="Неверный GUID 1С")
    conn = get_connection()
    try:
        begin_write(conn)
        obj = conn.execute("SELECT id, name, guid_1c FROM objects WHERE id = ?", (object_id,)).fetchone()
        if obj is None:
            raise HTTPException(status_code=404, detail="Объект не найден")
        if guid:
            other = conn.execute("SELECT name FROM objects WHERE guid_1c = ? AND id <> ?", (guid, object_id)).fetchone()
            if other:
                raise HTTPException(status_code=409, detail=f"Этот GUID уже привязан к объекту «{other['name']}»")
        conn.execute("UPDATE objects SET guid_1c = ?, updated_at = datetime('now') WHERE id = ?", (guid, object_id))
        conn.commit()
    finally:
        conn.close()
    activity.log("headcount_object_link", user=admin, entity_type="object", entity_id=object_id,
                 old_value=obj["guid_1c"], new_value=guid)
    return {"object_id": object_id, "guid_1c": guid}
