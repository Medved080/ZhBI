"""Производственная программа в деньгах: хранение, загрузка выгрузки, общие правки сроков, группы объектов, выгрузка в Excel
(2026-10-23, Docs/pp-money.md).

Расчёт сроков и распределения денег идёт в БРАУЗЕРЕ (`app/static/pp/core.js`): при перетаскивании полосы пересчёт нужен на каждое
движение мыши. Сервер хранит данные, проверяет права и отдаёт строки в том же формате, что `data.js` прототипа — `{money, schedule}`
с датами в серийных числах Excel, — поэтому ядро получает ровно те данные, на которых оно сверено с прототипом.

Права — два сквозных раздела (`app/features.py`): `pp` («Чтение» — просмотр и выгрузка, «Изменение» — общие правки сроков и группы)
и `pp_admin` (загрузка выгрузки). Заводски не выданы никому, кроме администратора сервиса. Раздел видит ВСЕ объекты программы, а не
только доступные пользователю объекты сервиса: деньги сводные, и право на раздел и есть доступ.

Загрузка ЗАМЕНЯЕТ работы и деньги целиком; объекты (по коду проекта), общие правки и группы сохраняются. Правка живёт по названию
работы, поэтому переживает перезагрузку; правка работы, которой в новой выгрузке нет, и объект, которого нет, — исчезают, и `analyze`
говорит об этом заранее. Объект ПП связывается с объектом сервиса по совпадению названия (иначе помечается «не в сервисе»).
"""

import io
import json
import math
import re
import sqlite3
from datetime import date, timedelta
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Body, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook

from app import activity
from app.access import require_service_feature
from app.auth import format_display_name
from app.backups import backup_before_import
from app.db import begin_write, get_connection
from app.pp_import import PpImportError, parse_workbook, short_smu
from app.upload_limits import read_upload_limited

router = APIRouter(prefix="/pp", tags=["pp"])

EPOCH = date(1970, 1, 1)
EXCEL_EPOCH_DAYS = 25569          # серийная дата Excel для 1970-01-01
MAX_EXPORT_ROWS = 300_000
MIN_DATE, MAX_DATE = date(2000, 1, 1), date(2100, 12, 31)


# ------------------------------------------------------------------ даты

def iso_of_serial(v) -> str:
    """Серийная дата Excel → ГГГГ-ММ-ДД. Округление как в ядре (`Math.round(v - 25569)`: половинки вверх)."""
    return (EPOCH + timedelta(days=math.floor(float(v) - EXCEL_EPOCH_DAYS + 0.5))).isoformat()


def serial_of_iso(s: str) -> int:
    return (date.fromisoformat(s) - EPOCH).days + EXCEL_EPOCH_DAYS


def _norm(s) -> str:
    return " ".join(str(s or "").split()).lower()


# ------------------------------------------------------------------ загрузка

def _service_objects(conn: sqlite3.Connection) -> dict:
    """Нормализованное название → id объекта сервиса (названия объектов уникальны)."""
    return {_norm(r["name"]): r["id"] for r in conn.execute("SELECT id, name FROM objects")}


def analyze_import(conn: sqlite3.Connection, parsed: dict) -> dict:
    """Что изменит загрузка — без записи: объекты (новые / исчезающие), работы, правки и группы, которые пропадут."""
    cur = {r["code"]: r["id"] for r in conn.execute("SELECT id, code FROM pp_objects")}
    new_codes = [o["code"] for o in parsed["objects"]]
    new_set = set(new_codes)
    removed = [c for c in cur if c not in new_set]
    added = [c for c in new_codes if c not in cur]
    works = {(r["c"], r["w"]) for r in parsed["schedule"]}
    orphan_edits = 0
    for r in conn.execute("SELECT o.code, e.work FROM pp_edits e JOIN pp_objects o ON o.id = e.pp_object_id"):
        if (r["code"], r["work"]) not in works:
            orphan_edits += 1
    members_lost = 0
    if removed:
        marks = ",".join("?" * len(removed))
        members_lost = conn.execute(
            f"SELECT COUNT(*) FROM pp_group_members m JOIN pp_objects o ON o.id = m.pp_object_id WHERE o.code IN ({marks})", removed
        ).fetchone()[0]
    index = _service_objects(conn)
    in_service = sum(1 for o in parsed["objects"] if _norm(o["name"]) in index)
    cur_ds = conn.execute("SELECT source_file, loaded_at FROM pp_dataset WHERE id = 1").fetchone()
    return {
        **parsed["stats"],
        "objects_new": len(added),
        "objects_removed": len(removed),
        "removed_codes": removed[:50],
        "edits_dropped": orphan_edits,
        "group_members_dropped": members_lost,
        "objects_in_service": in_service,
        "objects_not_in_service": len(parsed["objects"]) - in_service,
        "current": {"source_file": cur_ds["source_file"], "loaded_at": cur_ds["loaded_at"]} if cur_ds else None,
    }


def apply_import(conn: sqlite3.Connection, parsed: dict, filename: str, user: sqlite3.Row, *, dry_run: bool = False) -> dict:
    """Заменить работы и деньги целиком; объекты — по коду (id сохраняются, значит, правки и группы тоже)."""
    report = analyze_import(conn, parsed)
    begin_write(conn)
    try:
        cur = {r["code"]: r for r in conn.execute("SELECT id, code, object_id FROM pp_objects")}
        index = _service_objects(conn)
        ids: dict = {}
        for i, o in enumerate(parsed["objects"]):
            old = cur.get(o["code"])
            match = index.get(_norm(o["name"]))
            if old:
                # связь с объектом сервиса, поставленная раньше, не затираем; пустую пробуем найти по названию
                object_id = old["object_id"] if old["object_id"] is not None else match
                conn.execute("UPDATE pp_objects SET name = ?, unit = ?, division = ?, object_id = ?, sort_order = ? WHERE id = ?",
                             (o["name"], o["smu"], o["div"], object_id, i, old["id"]))
                ids[o["code"]] = old["id"]
            else:
                c = conn.execute("INSERT INTO pp_objects (code, name, unit, division, object_id, sort_order) VALUES (?, ?, ?, ?, ?, ?)",
                                 (o["code"], o["name"], o["smu"], o["div"], match, i))
                ids[o["code"]] = c.lastrowid
        gone = [c for c in cur if c not in ids]
        if gone:
            conn.executemany("DELETE FROM pp_objects WHERE code = ?", [(c,) for c in gone])   # каскадом уходят правки и участие в группах
        conn.execute("DELETE FROM pp_works")
        conn.execute("DELETE FROM pp_money")
        conn.executemany(
            "INSERT INTO pp_works (pp_object_id, work, start_date, end_date, parent_name, sort_order) VALUES (?, ?, ?, ?, ?, ?)",
            [(ids[r["c"]], r["w"], iso_of_serial(r["s"]), iso_of_serial(r["e"]), r["p"], i) for i, r in enumerate(parsed["schedule"])],
        )
        conn.executemany(
            "INSERT INTO pp_money (pp_object_id, work, price, done, rest, sort_order) VALUES (?, ?, ?, ?, ?, ?)",
            [(ids[r["c"]], r["w"], r["p"], r["dn"], r["r"], i) for i, r in enumerate(parsed["money"])],
        )
        conn.execute("DELETE FROM pp_edits WHERE NOT EXISTS (SELECT 1 FROM pp_works w WHERE w.pp_object_id = pp_edits.pp_object_id AND w.work = pp_edits.work)")
        conn.execute(
            "INSERT INTO pp_dataset (id, source_file, loaded_at, loaded_by, stats) VALUES (1, ?, datetime('now'), ?, ?) "
            "ON CONFLICT (id) DO UPDATE SET source_file = excluded.source_file, loaded_at = excluded.loaded_at, "
            "loaded_by = excluded.loaded_by, stats = excluded.stats",
            (filename, user["id"], json.dumps(parsed["stats"], ensure_ascii=False)),
        )
        if dry_run:
            conn.rollback()
        else:
            conn.commit()
    except Exception:
        conn.rollback()
        raise
    report["dry_run"] = dry_run
    return report


# ------------------------------------------------------------------ чтение

def read_raw(conn: sqlite3.Connection, ids: Optional[list] = None) -> dict:
    """Данные в формате `data.js` прототипа: {money: [{c,d,o,w,p,dn,r}], schedule: [{d,c,o,w,s,e,p}]}, порядок строк как в выгрузке."""
    where, args = "", []
    if ids:
        where = f"WHERE o.id IN ({','.join('?' * len(ids))})"
        args = list(ids)
    schedule = [
        {"d": r["unit"] or "", "c": r["code"], "o": r["name"], "w": r["work"], "s": serial_of_iso(r["start_date"]),
         "e": serial_of_iso(r["end_date"]), "p": r["parent_name"]}
        for r in conn.execute(
            f"SELECT o.code, o.name, o.unit, w.work, w.start_date, w.end_date, w.parent_name FROM pp_works w "
            f"JOIN pp_objects o ON o.id = w.pp_object_id {where} ORDER BY o.sort_order, o.id, w.sort_order, w.id", args)
    ]
    money = [
        {"c": r["code"], "d": short_smu(r["unit"] or ""), "o": r["name"], "w": r["work"], "p": r["price"], "dn": r["done"], "r": r["rest"]}
        for r in conn.execute(
            f"SELECT o.code, o.name, o.unit, m.work, m.price, m.done, m.rest FROM pp_money m "
            f"JOIN pp_objects o ON o.id = m.pp_object_id {where} ORDER BY o.sort_order, o.id, m.sort_order, m.id", args)
    ]
    return {"money": money, "schedule": schedule}


def read_meta(conn: sqlite3.Connection) -> dict:
    ds = conn.execute(
        "SELECT d.source_file, d.loaded_at, d.stats, u.last_name, u.first_name, u.patronymic FROM pp_dataset d "
        "LEFT JOIN users u ON u.id = d.loaded_by WHERE d.id = 1").fetchone()
    objects = [
        {"id": r["id"], "code": r["code"], "name": r["name"], "unit": r["unit"] or "", "object_id": r["object_id"],
         "in_service": r["object_id"] is not None}
        for r in conn.execute("SELECT id, code, name, unit, object_id FROM pp_objects ORDER BY sort_order, id")
    ]
    members: dict = {}
    for r in conn.execute("SELECT group_id, pp_object_id FROM pp_group_members ORDER BY group_id, pp_object_id"):
        members.setdefault(r["group_id"], []).append(r["pp_object_id"])
    groups = [{"id": r["id"], "name": r["name"], "ids": members.get(r["id"], [])}
              for r in sorted(conn.execute("SELECT id, name FROM pp_groups"), key=lambda r: r["name"].casefold())]
    return {
        "dataset": None if not ds else {
            "source_file": ds["source_file"], "loaded_at": ds["loaded_at"], "stats": json.loads(ds["stats"] or "{}"),
            "loaded_by": format_display_name(ds) or None,
        },
        "objects": objects,
        "groups": groups,
    }


def read_edits(conn: sqlite3.Connection) -> list:
    return [
        {"code": r["code"], "work": r["work"], "s": serial_of_iso(r["start_date"]), "e": serial_of_iso(r["end_date"]),
         "at": r["edited_at"], "by": format_display_name(r) or None}
        for r in conn.execute(
            "SELECT o.code, e.work, e.start_date, e.end_date, e.edited_at, u.last_name, u.first_name, u.patronymic FROM pp_edits e "
            "JOIN pp_objects o ON o.id = e.pp_object_id LEFT JOIN users u ON u.id = e.edited_by ORDER BY o.sort_order, e.work")
    ]


# ------------------------------------------------------------------ правки

def _check_dates(s, e) -> tuple:
    try:
        si, ei = iso_of_serial(s), iso_of_serial(e)
        ds, de = date.fromisoformat(si), date.fromisoformat(ei)
    except (TypeError, ValueError, OverflowError):
        raise HTTPException(status_code=400, detail="Неверная дата: ожидаются серийные числа Excel")
    if not (MIN_DATE <= ds <= MAX_DATE and MIN_DATE <= de <= MAX_DATE):
        raise HTTPException(status_code=400, detail="Дата вне допустимого диапазона 2000–2100 годов")
    if de < ds:
        raise HTTPException(status_code=400, detail="Окончание раньше начала")
    return si, ei


def save_edits(conn: sqlite3.Connection, user: sqlite3.Row, sets: list, deletes: list) -> dict:
    """Записать пачку правок листовых работ атомарно. Работа, которой нет в выгрузке, — ошибка всей пачки."""
    codes = {r["code"]: r["id"] for r in conn.execute("SELECT id, code FROM pp_objects")}
    begin_write(conn)
    try:
        saved = 0
        for item in sets:
            oid = codes.get(str(item.get("code", "")))
            work = str(item.get("work", ""))
            if oid is None or not conn.execute("SELECT 1 FROM pp_works WHERE pp_object_id = ? AND work = ?", (oid, work)).fetchone():
                raise HTTPException(status_code=400, detail="Правка относится к работе, которой нет в загруженной выгрузке")
            si, ei = _check_dates(item.get("s"), item.get("e"))
            conn.execute(
                "INSERT INTO pp_edits (pp_object_id, work, start_date, end_date, edited_by, edited_at) VALUES (?, ?, ?, ?, ?, datetime('now')) "
                "ON CONFLICT (pp_object_id, work) DO UPDATE SET start_date = excluded.start_date, end_date = excluded.end_date, "
                "edited_by = excluded.edited_by, edited_at = excluded.edited_at",
                (oid, work, si, ei, user["id"]))
            saved += 1
        removed = 0
        for item in deletes:
            oid = codes.get(str(item.get("code", "")))
            if oid is not None:
                removed += conn.execute("DELETE FROM pp_edits WHERE pp_object_id = ? AND work = ?", (oid, str(item.get("work", "")))).rowcount
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return {"saved": saved, "removed": removed}


# ------------------------------------------------------------------ группы

def _clean_group(name, ids, conn: sqlite3.Connection) -> tuple:
    name = " ".join(str(name or "").split())
    if not name:
        raise HTTPException(status_code=400, detail="Укажите название группы")
    if len(name) > 100:
        raise HTTPException(status_code=400, detail="Название группы длиннее 100 знаков")
    if not isinstance(ids, list):
        raise HTTPException(status_code=400, detail="Состав группы: ожидается список объектов")
    try:
        uniq = sorted({int(i) for i in ids})
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Состав группы: ожидаются номера объектов")
    if uniq:
        marks = ",".join("?" * len(uniq))
        found = conn.execute(f"SELECT COUNT(*) FROM pp_objects WHERE id IN ({marks})", uniq).fetchone()[0]
        if found != len(uniq):
            raise HTTPException(status_code=400, detail="В составе группы есть объекты, которых нет в программе")
    return name, uniq


def save_group(conn: sqlite3.Connection, user: sqlite3.Row, group_id: Optional[int], name, ids) -> int:
    name, ids = _clean_group(name, ids, conn)
    begin_write(conn)
    try:
        # регистр сравниваем в Python: NOCASE в SQLite не знает кириллицы
        clash = next((r for r in conn.execute("SELECT id, name FROM pp_groups") if r["name"].casefold() == name.casefold()), None)
        if clash and clash["id"] != group_id:
            raise HTTPException(status_code=409, detail=f"Группа «{name}» уже есть")
        if group_id is None:
            group_id = conn.execute("INSERT INTO pp_groups (name, created_by) VALUES (?, ?)", (name, user["id"])).lastrowid
        else:
            if not conn.execute("SELECT 1 FROM pp_groups WHERE id = ?", (group_id,)).fetchone():
                raise HTTPException(status_code=404, detail="Группа не найдена")
            conn.execute("UPDATE pp_groups SET name = ?, updated_at = datetime('now') WHERE id = ?", (name, group_id))
            conn.execute("DELETE FROM pp_group_members WHERE group_id = ?", (group_id,))
        conn.executemany("INSERT INTO pp_group_members (group_id, pp_object_id) VALUES (?, ?)", [(group_id, i) for i in ids])
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return group_id


# ------------------------------------------------------------------ выгрузка Excel

_NUM_FMT = "DD.MM.YYYY"


def build_export(money: list, schedule: list) -> bytes:
    """Книга из двух листов, как в прототипе: «Деньги» и «Сроки» (с колонкой «Правка»). Строки приходят от клиента — это его
    собственная выборка с пересчитанными сроками (`exportTables` в core.js), сервер только оформляет файл."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Деньги"
    ws.append(["Код проекта", "Дивизион", "Название объекта", "Наименование работ", "Договорная цена", "Выполнено", "Остаток"])
    for m in money:
        ws.append([m.get("c"), m.get("d"), m.get("o"), m.get("w"), m.get("p"), m.get("dn"), m.get("r")])
    ws2 = wb.create_sheet("Сроки")
    ws2.append(["Дивизион", "Код проекта", "Название объекта", "Наименование работ", "", "", "Начало", "Окончание", "Родитель", "Правка"])
    for i, s in enumerate(schedule, start=2):
        ws2.append([s.get("d"), s.get("c"), s.get("o"), s.get("w"), None, None, s.get("s"), s.get("e"), s.get("p"), s.get("mark") or None])
        ws2.cell(row=i, column=7).number_format = _NUM_FMT
        ws2.cell(row=i, column=8).number_format = _NUM_FMT
    for sheet, widths in ((ws, (16, 22, 40, 60, 18, 18, 18)), (ws2, (22, 16, 40, 60, 2, 2, 13, 13, 40, 18))):
        for col, w in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + col)].width = w
        sheet.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ API

def _guard(fn):
    try:
        return fn()
    except (PpImportError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc))


def _parse_upload(file: UploadFile) -> dict:
    payload = read_upload_limited(file.file)
    try:
        return _guard(lambda: parse_workbook(io.BytesIO(payload)))
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Файл не читается как Excel (.xlsx) с листами «Ввод» и «Лист6»")


@router.get("/meta")
def api_meta(user: sqlite3.Row = Depends(require_service_feature("pp", "read"))):
    conn = get_connection()
    try:
        meta = read_meta(conn)
        meta["edits_count"] = conn.execute("SELECT COUNT(*) FROM pp_edits").fetchone()[0]
        return meta
    finally:
        conn.close()


@router.get("/data")
def api_data(ids: Optional[str] = Query(None, description="id объектов ПП через запятую; без параметра — все"),
             user: sqlite3.Row = Depends(require_service_feature("pp", "read"))):
    try:
        wanted = [int(x) for x in ids.split(",") if x.strip()] if ids else None
    except ValueError:
        raise HTTPException(status_code=400, detail="ids: ожидаются номера объектов через запятую")
    conn = get_connection()
    try:
        return read_raw(conn, wanted)
    finally:
        conn.close()


@router.get("/edits")
def api_edits(user: sqlite3.Row = Depends(require_service_feature("pp", "read"))):
    conn = get_connection()
    try:
        return {"edits": read_edits(conn)}
    finally:
        conn.close()


@router.put("/edits")
def api_save_edits(body: dict = Body(...), user: sqlite3.Row = Depends(require_service_feature("pp", "write"))):
    """{"set": [{code, work, s, e}], "delete": [{code, work}]} — s и e серийными числами Excel."""
    sets, dels = body.get("set") or [], body.get("delete") or []
    if not isinstance(sets, list) or not isinstance(dels, list) or len(sets) + len(dels) > 20000:
        raise HTTPException(status_code=400, detail="Неверный состав пачки правок")
    conn = get_connection()
    try:
        result = save_edits(conn, user, sets, dels)
    finally:
        conn.close()
    activity.log("pp_edit", user=user, details={"set": result["saved"], "removed": result["removed"]})
    return result


@router.delete("/edits")
def api_reset_edits(codes: Optional[str] = Query(None, description="коды объектов через запятую; пусто — правки ВСЕХ объектов"),
                    user: sqlite3.Row = Depends(require_service_feature("pp", "write"))):
    conn = get_connection()
    try:
        begin_write(conn)
        try:
            if codes:
                lst = [c for c in (x.strip() for x in codes.split(",")) if c]
                marks = ",".join("?" * len(lst))
                n = conn.execute(f"DELETE FROM pp_edits WHERE pp_object_id IN (SELECT id FROM pp_objects WHERE code IN ({marks}))", lst).rowcount if lst else 0
            else:
                n = conn.execute("DELETE FROM pp_edits").rowcount
            conn.commit()
        except Exception:
            conn.rollback()
            raise
    finally:
        conn.close()
    activity.log("pp_edits_reset", user=user, new_value=(codes or "все объекты")[:200], details={"removed": n})
    return {"removed": n}


@router.post("/groups")
def api_group_create(body: dict = Body(...), user: sqlite3.Row = Depends(require_service_feature("pp", "write"))):
    conn = get_connection()
    try:
        gid = save_group(conn, user, None, body.get("name"), body.get("ids", []))
    finally:
        conn.close()
    activity.log("pp_group_save", user=user, entity_type="pp_group", entity_id=gid, new_value=str(body.get("name"))[:100])
    return {"id": gid}


@router.put("/groups/{group_id}")
def api_group_update(group_id: int, body: dict = Body(...), user: sqlite3.Row = Depends(require_service_feature("pp", "write"))):
    conn = get_connection()
    try:
        save_group(conn, user, group_id, body.get("name"), body.get("ids", []))
    finally:
        conn.close()
    activity.log("pp_group_save", user=user, entity_type="pp_group", entity_id=group_id, new_value=str(body.get("name"))[:100])
    return {"id": group_id}


@router.delete("/groups/{group_id}")
def api_group_delete(group_id: int, user: sqlite3.Row = Depends(require_service_feature("pp", "write"))):
    conn = get_connection()
    try:
        row = conn.execute("SELECT name FROM pp_groups WHERE id = ?", (group_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Группа не найдена")
        begin_write(conn)
        conn.execute("DELETE FROM pp_groups WHERE id = ?", (group_id,))
        conn.commit()
    finally:
        conn.close()
    activity.log("pp_group_delete", user=user, entity_type="pp_group", entity_id=group_id, old_value=row["name"])
    return {"id": group_id}


@router.post("/export")
def api_export(body: dict = Body(...), user: sqlite3.Row = Depends(require_service_feature("pp", "read"))):
    money, schedule = body.get("money") or [], body.get("schedule") or []
    if not isinstance(money, list) or not isinstance(schedule, list) or len(money) + len(schedule) > MAX_EXPORT_ROWS:
        raise HTTPException(status_code=400, detail="Неверный состав выгрузки")
    data = build_export(money, schedule)
    name = f"ПП_правки_{date.today().isoformat()}.xlsx"
    activity.log("pp_export", user=user, new_value=name, details={"money": len(money), "schedule": len(schedule)})
    return Response(
        data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


@router.post("/import/analyze")
def api_import_analyze(file: UploadFile = File(...), admin: sqlite3.Row = Depends(require_service_feature("pp_admin", "write"))):
    parsed = _parse_upload(file)
    conn = get_connection()
    try:
        return analyze_import(conn, parsed)
    finally:
        conn.close()


@router.post("/import/apply")
def api_import_apply(file: UploadFile = File(...), dry_run: bool = Form(False),
                     admin: sqlite3.Row = Depends(require_service_feature("pp_admin", "write"))):
    parsed = _parse_upload(file)
    filename = file.filename or "выгрузка.xlsx"
    if not dry_run:
        # копия базы ДО первой записи (правило проекта для любой загрузки из файла); не снялась — загрузка не стартует
        backup_before_import(f"производственная программа из {filename}", format_display_name(admin), admin["id"])
    conn = get_connection()
    try:
        result = apply_import(conn, parsed, filename, admin, dry_run=dry_run)
    finally:
        conn.close()
    if not dry_run:
        activity.log("pp_import", user=admin, new_value=filename,
                     details={k: result[k] for k in ("objects", "schedule_rows", "money_rows", "objects_removed", "edits_dropped")})
    return result
