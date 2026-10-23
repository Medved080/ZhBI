"""Производственная программа в деньгах: разбор исходной выгрузки `ДД.ММ.ГГГГ.xlsx` (2026-10-23, Docs/pp-money.md).

Это построчный перенос скрипта `transform.js` из прототипа «ПП 2026–2027» (он собирал `data.js` для дашборда). Правила не менялись,
поэтому результат обязан совпасть с прототипом: на файле 20.10.2026 — 213 объектов, 8 051 строка сроков, 5 449 строк денег,
остаток 442 284 млн ₽ (сверка — `scripts/verify_pp_import.py`). Модуль только РАЗБИРАЕТ файл и ничего не пишет в базу:
запись, привязка к объектам системы и права — следующие шаги.

Источник — книга с двумя листами:

* «Ввод» — НЕ таблица, а блоки: строка с кодом проекта и названием в колонке E открывает объект, ниже идут его работы со сроками
  (колонки H/I — начало и окончание СМР в серийных датах Excel). СМУ и дивизион в колонках A/B заполнены не в каждой строке и
  тянутся вниз от последней заполненной. Строки без обеих дат пропускаются.
* «Лист6» — деньги: строки по кодификатору. Суммируются по паре (код проекта, «Наименование работ по кодификатору», колонка R);
  родитель узла — колонка S; цена — J («Стоимость ОДЦ»), выполнено — L, остаток — M. Колонки «Невыбираемые» и знак остатка
  НЕ фильтруются: отрицательные остатки отбрасывает расчёт при показе (`app/static/pp/core.js`).

Иерархия сроков строится по НОМЕРУ в названии работы («3.2.1. …» → родитель — ближайшая выше работа «3.2.»), колонка
«Родитель по кодификатору» в «Ввод» не используется. Одинаковые названия внутри объекта (разные PBI-блоки: «Паркинг», «Надземная
часть») склеиваются в одну работу: начало — минимум, окончание — максимум, иначе деньги считались бы дважды.

Деньги объекта, у которых нет строки срока, получают сроки родителя с тем же номером раздела, а если такого нет — размах всего
объекта; если у объекта в «Ввод» вообще нет работ — запасной срок (`fallback_start`/`fallback_end`, в прототипе зашит как
01.01.2026–30.12.2027; здесь параметр). Деньги кодов, которых нет в «Ввод», в результат не попадают — их число сообщается в
`stats['codes_money_only']`.
"""

import math
import re
from datetime import date, datetime
from typing import IO, Optional, Union

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

SHEET_INPUT = "Ввод"
SHEET_MONEY = "Лист6"

# Серийные даты запасного срока объекта без работ в «Ввод»: 01.01.2026 и 30.12.2027 (значения прототипа).
FALLBACK_START = 46023
FALLBACK_END = 46751

_WK_PREFIX = re.compile(r"^\s*(\d+(?:\.\d+)*)\.")
_SMU_FULL = re.compile(r"^Строительно-монтажное управление-(\d+)$")


class PpImportError(ValueError):
    """Файл не похож на выгрузку производственной программы (нет листа, нет данных)."""


# Значения-ошибки Excel (ячейка с формулой, упавшей в #REF! и т. п.). SheetJS в прототипе читал их как ПУСТЫЕ ячейки — и название
# СМУ тянулось вниз от предыдущей заполненной строки; openpyxl отдаёт текст ошибки, его нужно сбросить тем же образом.
_EXCEL_ERRORS = frozenset({"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA"})


def _cell(row, i):
    v = row[i] if row is not None and i < len(row) else None
    return None if isinstance(v, str) and v.strip() in _EXCEL_ERRORS else v


def _js_string(v) -> str:
    """String(v) из JS: целое число без «.0», строка как есть."""
    if isinstance(v, float) and math.isfinite(v) and v.is_integer() and abs(v) < 1e21:
        return str(int(v))
    return str(v)


def _str(v) -> str:
    return "" if v is None else _js_string(v).strip()


def _num(v) -> float:
    """Число из ячейки; дата — серийным числом Excel; всё остальное (текст, пусто) — NaN, как в прототипе."""
    if isinstance(v, bool):
        return math.nan
    if isinstance(v, (datetime, date)):
        return float(to_excel(v))
    if isinstance(v, (int, float)) and math.isfinite(v):
        return float(v)
    return math.nan


def _norm_name(s) -> str:
    return re.sub(r"\s+", " ", str(s)).strip()


def _prefix(name) -> str:
    m = _WK_PREFIX.match(str(name))
    return m.group(1) if m else ""


def _level(prefix: str) -> int:
    return len(prefix.split("."))


def _short_smu(s: str) -> str:
    return _SMU_FULL.sub(r"СМУ-\1", s)


short_smu = _short_smu   # сокращение «Строительно-монтажное управление-N» → «СМУ-N» нужно и при выдаче денег (app/pp.py)


def _round2(x: float) -> float:
    """Math.round(x * 100) / 100 из JS: половинки округляются вверх (а не к чётному, как round())."""
    return math.floor(x * 100 + 0.5) / 100


def _whole(x: float):
    """Целые серийные даты — int (в JSON как в прототипе: 46300, а не 46300.0)."""
    return int(x) if float(x).is_integer() else x


def parse_workbook(source: Union[str, IO[bytes]], *, fallback_start: int = FALLBACK_START, fallback_end: int = FALLBACK_END) -> dict:
    """Разбирает книгу. Возвращает {'money': [...], 'schedule': [...], 'objects': [...], 'stats': {...}}.

    money — {c, d, o, w, p, dn, r}; schedule — {d, c, o, w, s, e, p}; формат `data.js` прототипа. objects — по объекту из «Ввод»:
    {code, name, smu, div} (дивизион прототип не использует, но источник его даёт).
    """
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        for name in (SHEET_INPUT, SHEET_MONEY):
            if name not in wb.sheetnames:
                raise PpImportError(f"В книге нет листа «{name}»")
        vvod = list(wb[SHEET_INPUT].iter_rows(values_only=True))
        list6 = list(wb[SHEET_MONEY].iter_rows(values_only=True))
    finally:
        wb.close()

    # ===== 1. «Ввод»: объекты, работы, подзадачи со сроками =====
    objects: dict[str, dict] = {}
    cur: Optional[dict] = None
    cur_smu = cur_div = ""
    rows: list[dict] = []
    for i in range(2, len(vvod)):
        r = vvod[i]
        if not r:
            continue
        smu, div, code = _str(_cell(r, 0)), _str(_cell(r, 1)), _str(_cell(r, 2))
        name5, name6 = _str(_cell(r, 4)), _str(_cell(r, 5))
        s, e = _num(_cell(r, 7)), _num(_cell(r, 8))
        if smu:
            cur_smu = smu
        if div:
            cur_div = div
        if code and name5:  # строка объекта
            cur = objects.get(code)
            if cur is None:
                cur = {"code": code, "name": name5, "div": cur_div or div, "smu": cur_smu or smu, "rows": [], "wks_by_num": {}}
                objects[code] = cur
            else:
                if cur_div or div:
                    cur["div"] = cur_div or div
                if cur_smu or smu:
                    cur["smu"] = cur_smu or smu
            rows = []
            continue
        if cur is None:
            continue
        nm = name5 or name6
        if not nm:
            continue
        if not (math.isfinite(s) and math.isfinite(e)):
            continue
        pfx = _prefix(nm)
        parent_name = None
        if pfx:
            lvl = _level(pfx)
            if lvl > 1:
                par_pfx = ".".join(pfx.split(".")[:-1])
                for prev in reversed(rows):
                    if prev["level"] == lvl - 1 and prev["pfx"] == par_pfx:
                        parent_name = prev["name"]
                        break
            n1 = int(pfx.split(".")[0])
            if n1 not in cur["wks_by_num"]:
                cur["wks_by_num"][n1] = {"name": nm, "s": s, "e": e}
        rows.append({"level": _level(pfx) if pfx else 0, "pfx": pfx, "name": nm})
        cur["rows"].append({"w": nm, "s": s, "e": e, "p": parent_name})
    if not objects:
        raise PpImportError("На листе «Ввод» не найдено ни одного объекта (строка с кодом проекта и названием)")

    # ===== 1б. Одинаковые работы в разных PBI-блоках → одна строка =====
    merged_dup = 0
    for o in objects.values():
        by_name: dict[str, dict] = {}
        merged: list[dict] = []
        for wk in o["rows"]:
            k = _norm_name(wk["w"]).lower()
            prev = by_name.get(k)
            if prev is not None:
                prev["s"] = min(prev["s"], wk["s"])
                prev["e"] = max(prev["e"], wk["e"])
                if not prev["p"] and wk["p"]:
                    prev["p"] = wk["p"]
                merged_dup += 1
            else:
                nw = {"w": _norm_name(wk["w"]), "s": wk["s"], "e": wk["e"], "p": wk["p"]}
                by_name[k] = nw
                merged.append(nw)
        o["rows"] = merged
        o["wks_by_num"] = {}
        for wk in o["rows"]:
            pfx = _prefix(wk["w"])
            if not pfx:
                continue
            n1 = int(pfx.split(".")[0])
            if n1 not in o["wks_by_num"]:
                o["wks_by_num"][n1] = {"name": wk["w"], "s": wk["s"], "e": wk["e"]}
    works_after_merge = sum(len(o["rows"]) for o in objects.values())

    # ===== 2. «Лист6»: деньги по (код, наименование работ по кодификатору) =====
    nodes_by_code: dict[str, dict[str, dict]] = {}
    for i in range(1, len(list6)):
        r = list6[i]
        if not r:
            continue
        code = _str(_cell(r, 0))
        if not code:
            continue
        name = _str(_cell(r, 17))
        if not name:
            continue
        m = nodes_by_code.setdefault(code, {})
        nd = m.get(name)
        if nd is None:
            nd = {"p": 0.0, "dn": 0.0, "r": 0.0, "par": _str(_cell(r, 18))}
            m[name] = nd
        for field, col in (("p", 9), ("dn", 11), ("r", 12)):
            v = _num(_cell(r, col))
            nd[field] += 0.0 if math.isnan(v) else v
    codes_money_only = sum(1 for code in nodes_by_code if code not in objects)

    # ===== 3. Сборка: «Сроки» (дерево) + «Деньги» =====
    sched: list[dict] = []
    money: list[dict] = []
    total_rest = 0.0
    skipped_empty = 0
    for o in objects.values():
        nodes = nodes_by_code.get(o["code"], {})
        emitted: set[str] = set()
        for wk in o["rows"]:
            sched.append({"d": o["smu"], "c": o["code"], "o": o["name"], "w": wk["w"], "s": _whole(wk["s"]), "e": _whole(wk["e"]), "p": wk["p"]})
            emitted.add(wk["w"])
        min_s, max_e = math.inf, -math.inf
        for wk in o["rows"]:
            min_s, max_e = min(min_s, wk["s"]), max(max_e, wk["e"])
        if not math.isfinite(min_s):
            if not nodes:
                skipped_empty += 1
                continue
            min_s, max_e = fallback_start, fallback_end
        synth: dict[str, None] = {}  # упорядоченное множество
        for nm, nd in nodes.items():
            if nd["par"] != nm and nd["par"] not in nodes:
                synth[nd["par"]] = None
        added: set[str] = set()

        def money_row(nm, nd, o=o):
            return {"c": o["code"], "d": _short_smu(o["smu"]), "o": o["name"], "w": nm,
                    "p": _round2(nd["p"]), "dn": _round2(nd["dn"]), "r": _round2(nd["r"])}

        def emit_node(nm, s, e, par, nodes=nodes, emitted=emitted, added=added, o=o):
            nonlocal total_rest
            if nm in added:
                return
            added.add(nm)
            nd = nodes.get(nm)
            if nm in emitted:
                # имя уже есть среди работ «Ввод» — добавляем только деньги, но обход не обрываем: у узла могут быть денежные дети
                if nd is not None:
                    money.append(money_row(nm, nd))
                    total_rest += nd["r"]
                for cnm, cnd in nodes.items():
                    if cnd["par"] == nm and cnm != nm:
                        emit_node(cnm, s, e, nm)
                return
            emitted.add(nm)
            sched.append({"d": o["smu"], "c": o["code"], "o": o["name"], "w": nm, "s": _whole(s), "e": _whole(e), "p": par or None})
            if nd is not None:
                money.append(money_row(nm, nd))
                total_rest += nd["r"]
            for cnm, cnd in nodes.items():
                if cnd["par"] == nm and cnm != nm:
                    emit_node(cnm, s, e, nm)

        def resolve_top(nm, o=o, min_s=min_s, max_e=max_e):
            pfx = _prefix(nm)
            w = o["wks_by_num"].get(int(pfx.split(".")[0])) if pfx else None
            return (w["name"], w["s"], w["e"]) if w else ("", min_s, max_e)

        for nm, nd in nodes.items():
            if nd["par"] == nm:
                p, s, e = resolve_top(nm)
                emit_node(nm, s, e, p)
        for sn in synth:
            p, s, e = resolve_top(sn)
            emit_node(sn, s, e, p)
        for nm, nd in nodes.items():
            if nm in added:
                continue
            if nd["par"] not in nodes and nd["par"] not in synth:
                p, s, e = resolve_top(nm)
                emit_node(nm, s, e, p)

    return {
        "money": money,
        "schedule": sched,
        "objects": [{"code": o["code"], "name": o["name"], "smu": o["smu"], "div": o["div"]} for o in objects.values()],
        "stats": {
            "objects": len(objects),
            "merged_duplicates": merged_dup,
            "works_after_merge": works_after_merge,
            "codes_with_money": len(nodes_by_code),
            "codes_money_only": codes_money_only,
            "schedule_rows": len(sched),
            "money_rows": len(money),
            "skipped_empty_objects": skipped_empty,
            "rest_total": total_rest,
        },
    }
