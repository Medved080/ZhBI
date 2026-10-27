"""
Универсальный классификатор видов работ (2026-10-27, живой запрос).

ОДИН справочник на все объекты (в отличие от `work_types`, который свой у
каждого объекта и грузился отдельным xlsx). Объект НАБИРАЕТ в свою структуру
работ конкретные позиции классификатора — набранная позиция становится
строкой `work_types` объекта со ссылкой `classifier_id`, поэтому блоки, факт,
сроки и отчёты работают с ней как раньше, без правок.

К виду работ привязывается любое число КАТЕГОРИЙ ЭЛЕМЕНТОВ («Монолит
стены», «СПК окна ПВХ»); к категории — любое число видов работ. Имена
категорий классификатора — свои, не имена категорий модели: с категориями
модели конкретного объекта (`revit_elements.category` — «Стены монолит»,
«Окна»…) их вручную сопоставляет человек, на объекте (`object_category_map`).
Через это сопоставление элемент модели находит вид работ, а по проценту
выполнения работы в его блоке рисуется прозрачным / с зелёными рёбрами /
обычным (`element_states`).

Формат файла — тот же, что у загрузчика видов работ (`app/work_types_import.py`:
лист `Works` или `WBS`, лист `PlanningTrack`), плюс лист `Elements`
(`Категория элементов | Название операции`): название операции в классификаторе
НЕ уникально («Устройство вертикальных ж/б конструкций» — ниже нуля и выше), и
категория привязывается ко ВСЕМ операциям с таким названием. Загрузка двухфазная
(`analyze` → `apply`), как у остальных справочников, и ТОЛЬКО ДОБАВЛЯЕТ связи
категорий: снять привязку можно лишь вручную, повторная загрузка файла чужую
правку не стирает.
"""

import io
import uuid

from openpyxl import load_workbook

from app import work_types_import as wti
from app.work_fact import current_percents_by_block_work

SHEET_ELEMENTS = "Elements"
ELEMENT_HEADERS = ("Категория элементов", "Название операции")

_PENDING = {}
_PENDING_LIMIT = 3

# Состояния элемента по работе в блоке (см. element_states). «Готово» и «работы
# нет» состоянием не называются: элемент рисуется как обычно.
STATE_NOT_STARTED = 1
STATE_IN_PROGRESS = 2


class ClassifierError(Exception):
    def __init__(self, status_code: int, message: str):
        self.status_code = status_code
        self.message = message
        super().__init__(message)


# ------------------------------------------------------------------ разбор файла

def parse_elements(data: bytes) -> tuple:
    """Лист `Elements` -> ([(категория, название операции)], предупреждения).
    Листа нет — пусто и без предупреждения: файл без привязок тоже годится."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ClassifierError(422, "Файл не открылся как xlsx: %s" % e)
    if SHEET_ELEMENTS not in wb.sheetnames:
        return [], []
    sheet = wb[SHEET_ELEMENTS]
    rows = sheet.iter_rows()
    try:
        header = next(rows)
    except StopIteration:
        return [], ["Лист «%s» пуст." % SHEET_ELEMENTS]
    try:
        idx = wti._index_by_headers(header, ELEMENT_HEADERS, (), SHEET_ELEMENTS)
    except wti.WorkTypesError as e:
        raise ClassifierError(e.status_code, e.message)
    pairs, warnings, seen = [], [], set()
    for n, row in enumerate(rows, start=2):
        def cell(col):
            i = idx[col]
            v = row[i].value if i < len(row) else None
            return str(v).strip() if v is not None else ""
        category, op = cell(ELEMENT_HEADERS[0]), cell(ELEMENT_HEADERS[1])
        if not category and not op:
            continue
        if not category or not op:
            warnings.append("Лист «%s», строка %d: заполнена только одна колонка, строка пропущена"
                            % (SHEET_ELEMENTS, n))
            continue
        if (category.casefold(), op) in seen:
            continue
        seen.add((category.casefold(), op))
        pairs.append((category, op))
    return pairs, warnings


# ---------------------------------------------------------------------- загрузка

def _existing(conn) -> dict:
    return {r["path"]: dict(r) for r in conn.execute(
        "SELECT id, path, row_kind, name, retired_at FROM work_classifier")}


def _ops_by_name(parsed) -> dict:
    result = {}
    for r in parsed:
        if r["row_kind"] in (wti.ROW_OP, wti.ROW_MILESTONE) and r["name"]:
            result.setdefault(r["name"], []).append(r["path"])
    return result


def analyze(conn, data: bytes) -> dict:
    """Фаза 1. В БД не пишет."""
    try:
        parsed, warnings, tracks = wti.parse_xlsx(data)
    except wti.WorkTypesError as e:
        raise ClassifierError(e.status_code, e.message)
    pairs, el_warnings = parse_elements(data)
    warnings = list(warnings) + el_warnings
    if not parsed:
        raise ClassifierError(422, "В файле нет ни одной строки видов работ.")

    have = _existing(conn)
    parsed_paths = {r["path"] for r in parsed}
    new_rows = [r for r in parsed if r["path"] not in have]
    reviving = [r for r in parsed if r["path"] in have and have[r["path"]]["retired_at"]]
    retiring = [p for p, row in have.items() if p not in parsed_paths and not row["retired_at"]]

    by_name = _ops_by_name(parsed)
    have_cats = {r["name"].casefold(): r["id"] for r in conn.execute(
        "SELECT id, name FROM work_element_categories")}
    categories = {}
    for category, op in pairs:
        entry = categories.setdefault(category.casefold(), {
            "категория": category, "новая": category.casefold() not in have_cats, "операции": []})
        paths = by_name.get(op, [])
        if not paths:
            warnings.append("Лист «%s»: операция «%s» (категория «%s») не найдена на листе видов работ"
                            % (SHEET_ELEMENTS, op, category))
        entry["операции"].extend(paths)

    return {
        "total_rows": len(parsed),
        "new": [{"путь": r["path"], "тип": r["row_kind"], "единица": r["unit"]} for r in new_rows],
        "reviving": [{"путь": r["path"]} for r in reviving],
        "retiring": sorted(retiring),
        "unchanged": len(parsed) - len(new_rows) - len(reviving),
        "tracks": [{"код": t["code"], "название": t["name"]} for t in tracks],
        "categories": [{"категория": c["категория"], "новая": c["новая"],
                        "операций": len(c["операции"])} for c in categories.values()],
        "warnings": warnings,
        "_parsed": parsed, "_tracks": tracks, "_categories": categories,
    }


def apply(conn, analysis: dict) -> dict:
    """Фаза 2. Сопоставление — по ПУТИ; пропавшие пути помечаются `retired_at`."""
    parsed, tracks = analysis["_parsed"], analysis["_tracks"]
    have = _existing(conn)
    path_to_id = {p: row["id"] for p, row in have.items()}
    added = revived = 0
    for r in parsed:
        existing = have.get(r["path"])
        parent_id = path_to_id.get(r["parent_path"]) if r["parent_path"] else None
        if existing is None:
            cur = conn.execute(
                "INSERT INTO work_classifier (parent_id, path, row_kind, code, name, unit, note, "
                "planning_track_code, sort_order) VALUES (?,?,?,?,?,?,?,?,?)",
                (parent_id, r["path"], r["row_kind"], r["code"], r["name"], r["unit"], r["note"],
                 r["track_code"], r["sort_order"]))
            path_to_id[r["path"]] = cur.lastrowid
            added += 1
        else:
            if existing["retired_at"]:
                revived += 1
            conn.execute(
                "UPDATE work_classifier SET parent_id = ?, code = ?, name = ?, unit = ?, note = ?, "
                "planning_track_code = ?, sort_order = ?, row_kind = ?, retired_at = NULL WHERE id = ?",
                (parent_id, r["code"], r["name"], r["unit"], r["note"], r["track_code"],
                 r["sort_order"], r["row_kind"], existing["id"]))
    parsed_paths = {r["path"] for r in parsed}
    retired = 0
    for path, row in have.items():
        if path not in parsed_paths and not row["retired_at"]:
            conn.execute("UPDATE work_classifier SET retired_at = datetime('now') WHERE id = ?", (row["id"],))
            retired += 1

    if tracks:
        conn.execute("DELETE FROM work_classifier_tracks")
        for t in tracks:
            conn.execute("INSERT INTO work_classifier_tracks (code, name, note) VALUES (?,?,?)",
                         (t["code"], t["name"], t["note"]))

    links = 0
    for entry in analysis["_categories"].values():
        category_id = _category_id(conn, entry["категория"], create=True)
        for path in entry["операции"]:
            cur = conn.execute(
                "INSERT OR IGNORE INTO work_classifier_categories (classifier_id, category_id) VALUES (?,?)",
                (path_to_id[path], category_id))
            links += cur.rowcount

    linked_objects = link_all_objects(conn)
    conn.commit()
    return {"added": added, "revived": revived, "retired": retired, "total": len(parsed),
            "categories": len(analysis["_categories"]), "links": links,
            "linked_positions": linked_objects}


def remember_pending(analysis: dict) -> str:
    token = uuid.uuid4().hex
    _PENDING[token] = analysis
    while len(_PENDING) > _PENDING_LIMIT:
        _PENDING.pop(next(iter(_PENDING)))
    return token


def get_pending(token: str) -> dict:
    analysis = _PENDING.get(token)
    if analysis is None:
        raise ClassifierError(410, "Результат разбора уже недоступен (сервер перезапускался или "
                                   "разбор устарел). Загрузите файл заново.")
    return analysis


def forget_pending(token: str) -> None:
    _PENDING.pop(token, None)


# ------------------------------------------------------------- чтение и правка

def _category_id(conn, name: str, create: bool = False):
    row = conn.execute("SELECT id FROM work_element_categories WHERE name = ?", (name,)).fetchone()
    if row:
        return row["id"]
    if not create:
        return None
    return conn.execute("INSERT INTO work_element_categories (name) VALUES (?)", (name,)).lastrowid


def tree(conn) -> dict:
    """Всё дерево классификатора со списком категорий у каждой позиции."""
    cats = {}
    for r in conn.execute(
            "SELECT l.classifier_id, c.id, c.name FROM work_classifier_categories l "
            "JOIN work_element_categories c ON c.id = l.category_id ORDER BY c.name"):
        cats.setdefault(r["classifier_id"], []).append({"id": r["id"], "name": r["name"]})
    rows = []
    for r in conn.execute(
            "SELECT id, parent_id, path, row_kind, code, name, unit, note, planning_track_code, retired_at "
            "FROM work_classifier WHERE retired_at IS NULL ORDER BY sort_order, id"):
        d = dict(r)
        d["categories"] = cats.get(r["id"], [])
        rows.append(d)
    tracks = {r["code"]: r["name"] for r in conn.execute("SELECT code, name FROM work_classifier_tracks")}
    categories = [dict(r) for r in conn.execute(
        "SELECT c.id, c.name, COUNT(l.classifier_id) AS works FROM work_element_categories c "
        "LEFT JOIN work_classifier_categories l ON l.category_id = c.id GROUP BY c.id ORDER BY c.name")]
    return {"rows": rows, "tracks": tracks, "categories": categories}


def set_position_categories(conn, classifier_id: int, names) -> dict:
    """Задать категории позиции ПОЛНЫМ списком имён (новые заводятся)."""
    row = conn.execute("SELECT row_kind FROM work_classifier WHERE id = ? AND retired_at IS NULL",
                       (classifier_id,)).fetchone()
    if row is None:
        raise ClassifierError(404, "Позиция классификатора не найдена")
    if row["row_kind"] == wti.ROW_NODE:
        raise ClassifierError(422, "Категории привязываются к операции или вехе, а не к разделу")
    if not isinstance(names, list):
        raise ClassifierError(422, "Список категорий некорректен")
    clean = []
    for n in names:
        n = (n or "").strip() if isinstance(n, str) else ""
        if not n:
            continue
        if len(n) > 120:
            raise ClassifierError(422, "Название категории длиннее 120 символов")
        if n.casefold() not in {c.casefold() for c in clean}:
            clean.append(n)
    ids = {_category_id(conn, n, create=True) for n in clean}
    conn.execute("DELETE FROM work_classifier_categories WHERE classifier_id = ?", (classifier_id,))
    conn.executemany("INSERT INTO work_classifier_categories (classifier_id, category_id) VALUES (?,?)",
                     [(classifier_id, i) for i in ids])
    conn.commit()
    return {"categories": len(ids)}


def delete_category(conn, category_id: int) -> dict:
    row = conn.execute("SELECT name FROM work_element_categories WHERE id = ?", (category_id,)).fetchone()
    if row is None:
        raise ClassifierError(404, "Категория не найдена")
    maps = conn.execute("SELECT COUNT(*) FROM object_category_map WHERE category_id = ?",
                        (category_id,)).fetchone()[0]
    conn.execute("DELETE FROM work_element_categories WHERE id = ?", (category_id,))
    conn.commit()
    return {"name": row["name"], "maps_removed": maps}


# ------------------------------------------------------- связь с работами объекта

def link_all_objects(conn) -> int:
    """Проставить `work_types.classifier_id` по совпадению пути — всем объектам.
    Объект, чей справочник загружен тем же файлом, подхватывает классификатор
    без ручной работы; строки с ДРУГИМ путём остаются своими."""
    cur = conn.execute(
        "UPDATE work_types SET classifier_id = (SELECT c.id FROM work_classifier c "
        "WHERE c.path = work_types.path AND c.retired_at IS NULL) "
        "WHERE classifier_id IS NULL AND EXISTS (SELECT 1 FROM work_classifier c "
        "WHERE c.path = work_types.path AND c.retired_at IS NULL)")
    return cur.rowcount


def link_object(conn, object_id: int) -> int:
    cur = conn.execute(
        "UPDATE work_types SET classifier_id = (SELECT c.id FROM work_classifier c "
        "WHERE c.path = work_types.path AND c.retired_at IS NULL) "
        "WHERE object_id = ? AND classifier_id IS NULL AND EXISTS (SELECT 1 FROM work_classifier c "
        "WHERE c.path = work_types.path AND c.retired_at IS NULL)", (object_id,))
    return cur.rowcount


def _object_active_paths(conn, object_id: int) -> dict:
    return {r["path"]: r["id"] for r in conn.execute(
        "SELECT id, path FROM work_types WHERE object_id = ? AND retired_at IS NULL", (object_id,))}


def object_selection(conn, object_id: int) -> dict:
    """Дерево классификатора + какие позиции уже в структуре работ объекта."""
    data = tree(conn)
    active = _object_active_paths(conn, object_id)
    selected = [r["id"] for r in data["rows"]
                if r["row_kind"] != wti.ROW_NODE and r["path"] in active]
    own = conn.execute(
        "SELECT COUNT(*) FROM work_types w WHERE w.object_id = ? AND w.retired_at IS NULL "
        "AND w.row_kind != 'узел' AND NOT EXISTS (SELECT 1 FROM work_classifier c "
        "WHERE c.path = w.path AND c.retired_at IS NULL)", (object_id,)).fetchone()[0]
    return {"rows": data["rows"], "tracks": data["tracks"], "selected": selected, "own_positions": own}


def _ancestors(rows_by_id: dict, leaf_id: int):
    chain, cur = [], rows_by_id[leaf_id]
    while cur is not None:
        chain.append(cur["id"])
        cur = rows_by_id.get(cur["parent_id"])
    return chain


def save_object_selection(conn, object_id: int, leaf_ids) -> dict:
    """Структура работ объекта = выбранные позиции + их разделы-предки.
    Недостающее добавляется, лишнее из набранного ранее — мягко списывается
    (`retired_at`, статусы и проценты не теряются). Строки объекта, которых в
    классификаторе нет, не трогаются."""
    if not isinstance(leaf_ids, list) or len(leaf_ids) != len(set(leaf_ids)):
        raise ClassifierError(422, "Список позиций некорректен")
    rows = {r["id"]: dict(r) for r in conn.execute(
        "SELECT id, parent_id, path, row_kind, code, name, unit, note, planning_track_code, sort_order "
        "FROM work_classifier WHERE retired_at IS NULL")}
    for i in leaf_ids:
        if i not in rows or rows[i]["row_kind"] == wti.ROW_NODE:
            raise ClassifierError(422, "В списке есть позиция, которой нет в классификаторе, или раздел")
    desired = set()
    for i in leaf_ids:
        desired.update(_ancestors(rows, i))
    order = sorted(desired, key=lambda i: (rows[i]["sort_order"], i))

    have = {r["path"]: dict(r) for r in conn.execute(
        "SELECT id, path, retired_at, classifier_id FROM work_types WHERE object_id = ?", (object_id,))}
    path_to_wt = {p: r["id"] for p, r in have.items()}
    added = revived = 0
    for i in order:
        c = rows[i]
        parent = rows.get(c["parent_id"])
        parent_id = path_to_wt.get(parent["path"]) if parent else None
        row = have.get(c["path"])
        if row is None:
            cur = conn.execute(
                "INSERT INTO work_types (object_id, parent_id, path, row_kind, code, name, unit, note, "
                "planning_track_code, sort_order, classifier_id) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (object_id, parent_id, c["path"], c["row_kind"], c["code"], c["name"], c["unit"],
                 c["note"], c["planning_track_code"], c["sort_order"], i))
            path_to_wt[c["path"]] = cur.lastrowid
            added += 1
        else:
            if row["retired_at"]:
                revived += 1
            conn.execute("UPDATE work_types SET classifier_id = ?, retired_at = NULL WHERE id = ?",
                         (i, row["id"]))

    desired_paths = {rows[i]["path"] for i in desired}
    class_paths = {r["path"] for r in conn.execute(
        "SELECT path FROM work_classifier WHERE retired_at IS NULL")}
    retired = used = 0
    for path, row in have.items():
        if row["retired_at"] or path in desired_paths:
            continue
        if row["classifier_id"] is None and path not in class_paths:
            continue   # своя строка объекта — классификатор её не касается
        in_use = conn.execute(
            "SELECT (SELECT COUNT(*) FROM block_works WHERE work_type_id = ? AND retired_at IS NULL) + "
            "(SELECT COUNT(*) FROM object_works WHERE work_type_id = ? AND retired_at IS NULL)",
            (row["id"], row["id"])).fetchone()[0]
        used += 1 if in_use else 0
        conn.execute("UPDATE work_types SET retired_at = datetime('now') WHERE id = ?", (row["id"],))
        retired += 1
    conn.commit()
    return {"added": added, "revived": revived, "retired": retired, "retired_in_use": used,
            "selected": len(leaf_ids)}


def preview_object_selection(conn, object_id: int, leaf_ids) -> dict:
    """Что изменит сохранение — без записи (для подтверждения в форме)."""
    if not isinstance(leaf_ids, list):
        raise ClassifierError(422, "Список позиций некорректен")
    rows = {r["id"]: dict(r) for r in conn.execute(
        "SELECT id, parent_id, path, row_kind FROM work_classifier WHERE retired_at IS NULL")}
    desired = set()
    for i in leaf_ids:
        if i in rows:
            desired.update(_ancestors(rows, i))
    desired_paths = {rows[i]["path"] for i in desired}
    class_paths = {r["path"] for r in rows.values()}
    leaves = {r["path"] for r in rows.values() if r["row_kind"] != wti.ROW_NODE}
    active = conn.execute(
        "SELECT id, path, classifier_id FROM work_types WHERE object_id = ? AND retired_at IS NULL "
        "AND row_kind != 'узел'", (object_id,)).fetchall()
    active_paths = {r["path"] for r in active}
    removing = [r for r in active if r["path"] not in desired_paths
                and (r["classifier_id"] is not None or r["path"] in class_paths)]
    in_use = 0
    for r in removing:
        in_use += 1 if conn.execute(
            "SELECT (SELECT COUNT(*) FROM block_works WHERE work_type_id = ? AND retired_at IS NULL) + "
            "(SELECT COUNT(*) FROM object_works WHERE work_type_id = ? AND retired_at IS NULL)",
            (r["id"], r["id"])).fetchone()[0] else 0
    adding = len({rows[i]["path"] for i in leaf_ids if i in rows and rows[i]["path"] in leaves} - active_paths)
    return {"adding": adding, "removing": len(removing), "removing_in_use": in_use}


# --------------------------------------------------- сопоставление категорий модели

def category_map_state(conn, object_id: int) -> dict:
    """Категории классификатора, действующие на объекте (у набранных позиций),
    категории его модели и текущее ручное сопоставление."""
    link_object(conn, object_id)
    conn.commit()
    categories = [dict(r) for r in conn.execute(
        "SELECT c.id, c.name, COUNT(DISTINCT w.id) AS works FROM work_element_categories c "
        "JOIN work_classifier_categories l ON l.category_id = c.id "
        "JOIN work_types w ON w.classifier_id = l.classifier_id AND w.object_id = ? AND w.retired_at IS NULL "
        "GROUP BY c.id ORDER BY c.name", (object_id,))]
    model = [dict(r) for r in conn.execute(
        "SELECT category AS name, COUNT(*) AS n FROM revit_elements "
        "WHERE object_id = ? AND is_current = 1 AND category IS NOT NULL AND category != '' "
        "GROUP BY category ORDER BY n DESC", (object_id,))]
    mapping = {}
    for r in conn.execute("SELECT category_id, model_category FROM object_category_map WHERE object_id = ?",
                          (object_id,)):
        mapping.setdefault(r["category_id"], []).append(r["model_category"])
    return {"categories": categories, "model_categories": model, "map": mapping}


def save_category_map(conn, object_id: int, mapping) -> dict:
    """Полная замена сопоставления объекта: {id категории: [категории модели]}."""
    if not isinstance(mapping, dict):
        raise ClassifierError(422, "Сопоставление некорректно")
    state = category_map_state(conn, object_id)
    valid_cat = {c["id"] for c in state["categories"]}
    valid_model = {m["name"] for m in state["model_categories"]}
    pairs = []
    for key, names in mapping.items():
        try:
            cid = int(key)
        except (TypeError, ValueError):
            raise ClassifierError(422, "Сопоставление некорректно")
        if cid not in valid_cat:
            raise ClassifierError(422, "Категория не действует на этом объекте")
        if not isinstance(names, list):
            raise ClassifierError(422, "Сопоставление некорректно")
        for n in names:
            if n not in valid_model:
                raise ClassifierError(422, "Категории «%s» нет в модели объекта" % n)
            pairs.append((object_id, cid, n))
    conn.execute("DELETE FROM object_category_map WHERE object_id = ?", (object_id,))
    conn.executemany("INSERT OR IGNORE INTO object_category_map (object_id, category_id, model_category) "
                     "VALUES (?,?,?)", pairs)
    conn.commit()
    return {"pairs": len(set(pairs))}


# ---------------------------------------------------------- состояние элементов

def element_states(conn, object_id: int) -> dict:
    """{(section_id, level_id): {категория модели: STATE_*}} — для отрисовки.

    Элемент находит категорию классификатора через ручное сопоставление
    объекта, а ту — виды работ ЭТОГО объекта. Берутся работы, которые есть в
    плане именно этого блока (`block_works`): нет работы в блоке — нет и
    состояния, элемент рисуется как обычно. Несколько работ: «начато» — если
    начата любая, «готово» (нет состояния) — только когда выполнены все.
    Возвращаются лишь состояния «не начато» и «в работе»."""
    pairs = conn.execute(
        "SELECT m.model_category, w.id AS work_type_id "
        "FROM object_category_map m "
        "JOIN work_classifier_categories l ON l.category_id = m.category_id "
        "JOIN work_types w ON w.classifier_id = l.classifier_id AND w.object_id = m.object_id "
        "AND w.retired_at IS NULL WHERE m.object_id = ?", (object_id,)).fetchall()
    if not pairs:
        return {}
    wt_to_cats = {}
    for r in pairs:
        wt_to_cats.setdefault(r["work_type_id"], set()).add(r["model_category"])

    percents = current_percents_by_block_work(conn, object_id)
    blocks = {r["id"]: (r["section_id"], r["level_id"]) for r in conn.execute(
        "SELECT id, section_id, level_id FROM blocks WHERE object_id = ?", (object_id,))}
    # (блок, категория модели) -> проценты всех действующих там работ
    acc = {}
    for r in conn.execute(
            "SELECT id, block_id, work_type_id FROM block_works "
            "WHERE object_id = ? AND retired_at IS NULL", (object_id,)):
        cats = wt_to_cats.get(r["work_type_id"])
        if not cats or r["block_id"] not in blocks:
            continue
        for cat in cats:
            acc.setdefault((r["block_id"], cat), []).append(percents.get(r["id"], 0))

    result = {}
    for (block_id, cat), values in acc.items():
        if all(v >= 100 for v in values):
            continue
        state = STATE_IN_PROGRESS if any(v > 0 for v in values) else STATE_NOT_STARTED
        result.setdefault(blocks[block_id], {})[cat] = state
    return result
