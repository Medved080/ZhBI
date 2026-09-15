"""
Плоская «Шахматка» (2026-09-14, Docs/design/chess-flat/CLAUDE-PROMPT.md) —
развёртка здания по этажам и секциям для просмотра, массового ввода факта и
печати бланка обхода, поверх уже существующей 2D/3D «Шахматки»
(app/work_fact.py). Геометрию и проценты этот модуль не пересчитывает
заново — источник истины прежний (`app/blocks.py`, `app/work_fact.py`);
здесь только раскладка готовых данных под плоский экран и атомарная запись
пакета из нескольких блоков разом, которой у одиночного
POST .../fact-reports нет.
"""

import json
import sqlite3

from app import blocks as blocks_mod
from app import work_fact
from app.work_fact import FactError
# Приватное имя по месту объявления (`app/work_fact.py`), но «текущий
# процент операции на блоке» — один и тот же расчёт что там, что здесь:
# проверка конфликта перед пакетной записью должна сравнивать ровно с тем
# значением, что показал экран, — вторая реализация разошлась бы при первой
# же правке правила «максимальная дата, при равенстве — больший id» (тот же
# приём, что у app/block_works.py).
from app.work_fact import _current_percents


class ConflictError(Exception):
    """Кто-то другой изменил факт после того, как эти данные показал
    экран (§7.8 задания) — пакет не пишется целиком, ни одной строкой."""

    def __init__(self, conflicts: list):
        self.conflicts = conflicts
        super().__init__("Данные на сервере изменились с момента ввода.")


def layout(conn: sqlite3.Connection, object_id: int, track_code: str) -> dict:
    """Всё, что нужно плоскому экрану разом: справочник операций доски,
    секции и уровни объекта (в порядке отображения) и по каждому
    существующему блоку — применимые операции этой доски с их текущим
    процентом (остальные операции доски у блока — «не применяется», см.
    `work_fact.board_block_values`)."""
    ops = work_fact.board_ops(conn, object_id, track_code)
    if not ops:
        raise FactError(404, "Доска «%s» не найдена или в ней нет операций." % track_code)
    sections = blocks_mod.list_sections(conn, object_id)
    # Уровни объекта хранятся по возрастанию отметки/номера (sort_order) —
    # развёртке нужен обратный порядок, сверху вниз (см. проверку на
    # обезличенной копии: кровля/верхние этажи — наибольший sort_order).
    levels = list(reversed(blocks_mod.list_levels(conn, object_id)))
    geometry = blocks_mod.list_blocks(conn, object_id)
    board = work_fact.board_block_values(conn, object_id, track_code)

    out_blocks = []
    for b in geometry:
        entry = board.get(b["id"])
        percents = {op["id"]: op["percent"] for op in entry["ops"]} if entry else {}
        out_blocks.append({
            "id": b["id"], "section_id": b["section_id"], "level_id": b["level_id"],
            "percents": percents,
        })
    obj = conn.execute("SELECT name FROM objects WHERE id = ?", (object_id,)).fetchone()
    return {
        "object_name": obj["name"] if obj else "",
        "ops": ops,
        "sections": [
            {"id": s["id"], "code": s["code"], "name": s["name"], "sort_order": s["sort_order"]}
            for s in sections
        ],
        "levels": [
            {"id": l["id"], "floor": l["floor"], "kind": l["kind"], "name": l["name"],
             "sort_order": l["sort_order"]}
            for l in levels
        ],
        "blocks": out_blocks,
    }


def commit_batch(conn: sqlite3.Connection, object_id: int, user_id: int, *,
                 track_code: str, report_date: str, idempotency_key: str,
                 items: list) -> dict:
    """Пакетная запись факта плоской «Шахматки» — весь пакет атомарно, или
    ни одной новой/дополненной строки (§7.6 задания). Каждый затронутый блок
    получает СВОЙ документ на `report_date` (новый или уже существующий на
    эту дату): в него добавляются/правятся ТОЛЬКО введённые операции — уже
    записанные в этом документе строки ДРУГИХ операций (в т.ч. с других
    досок) сохраняются как есть, а не переписываются скрытым текущим
    значением (§7.5 — иначе полный слепок `work_fact.save_report`
    незаметно стёр бы соседние операции того же отчёта)."""
    if not report_date:
        raise FactError(422, "Не указана дата факта.")
    if not items:
        raise FactError(422, "Нет заполненных значений для записи.")
    if not idempotency_key:
        raise FactError(422, "Не передан ключ идемпотентности.")

    # Повтор того же запроса (таймаут, повторная отправка браузером) —
    # отдаём уже сохранённый результат, вторая транзакция не открывается.
    existing_batch = conn.execute(
        "SELECT result_json FROM chess_flat_batches WHERE object_id = ? AND idempotency_key = ?",
        (object_id, idempotency_key),
    ).fetchone()
    if existing_batch:
        return json.loads(existing_batch["result_json"])

    board_op_ids = {o["id"] for o in work_fact.board_ops(conn, object_id, track_code)}
    if not board_op_ids:
        raise FactError(404, "Доска «%s» не найдена или в ней нет операций." % track_code)

    block_ids = sorted({it["block_id"] for it in items})
    placeholders = ",".join("?" * len(block_ids))
    real_blocks = {
        r["id"] for r in conn.execute(
            "SELECT id FROM blocks WHERE object_id = ? AND id IN (%s)" % placeholders,
            (object_id, *block_ids),
        )
    }
    missing = set(block_ids) - real_blocks
    if missing:
        raise FactError(404, "Блоки не найдены на объекте: %s." % sorted(missing))

    selected_by_block: dict = {}
    current_by_block: dict = {}
    by_block: dict = {}
    conflicts = []
    for it in items:
        block_id, wt_id, percent = it["block_id"], it["work_type_id"], it["percent"]
        if not isinstance(percent, int) or not (0 <= percent <= 100):
            raise FactError(422, "Процент вне 0..100 у блока %s, операции %s." % (block_id, wt_id))
        if wt_id not in board_op_ids:
            raise FactError(422, "Операция %s не входит в доску «%s»." % (wt_id, track_code))
        if block_id not in selected_by_block:
            selected_by_block[block_id] = set(
                work_fact.block_settings(conn, object_id, block_id)["selected"])
            current_by_block[block_id] = _current_percents(conn, block_id)
        if wt_id not in selected_by_block[block_id]:
            raise FactError(422, "Операция %s не применяется к блоку %s." % (wt_id, block_id))
        actual = current_by_block[block_id].get(wt_id, 0)
        expected = it.get("expected_percent", 0)
        if actual != expected:
            conflicts.append({
                "block_id": block_id, "work_type_id": wt_id,
                "expected": expected, "actual": actual,
            })
        by_block.setdefault(block_id, {})[wt_id] = percent

    if conflicts:
        raise ConflictError(conflicts)

    # Ни одного conn.commit() до этой точки (save_report вызывается с
    # commit=False) — пакет копится в ОДНОЙ транзакции этого соединения,
    # атомарность обеспечивает единственный commit ниже.
    report_ids = {}
    for block_id, touched in by_block.items():
        existing_report = conn.execute(
            "SELECT id FROM work_fact_reports WHERE object_id = ? AND block_id = ? "
            "AND report_date = ?",
            (object_id, block_id, report_date),
        ).fetchone()
        old_items = {}
        if existing_report:
            old_items = work_fact.get_report(conn, object_id, block_id, existing_report["id"])["items"]
        merged = {**old_items, **touched}
        report_id = work_fact.save_report(
            conn, object_id, user_id, block_id,
            existing_report["id"] if existing_report else None,
            report_date, merged, commit=False,
        )
        report_ids[block_id] = report_id

    result = {
        "report_date": report_date,
        "items_count": len(items),
        "blocks_count": len(by_block),
        "reports": report_ids,
    }
    conn.execute(
        "INSERT INTO chess_flat_batches (object_id, idempotency_key, track_code, report_date, "
        "created_by, result_json) VALUES (?,?,?,?,?,?)",
        (object_id, idempotency_key, track_code, report_date, user_id, json.dumps(result)),
    )
    conn.commit()
    return result


# ================ Бланк обхода: выгрузка в PDF/XLSX (2026-09-15) ================
#
# Строки уже посчитаны и сгруппированы на клиенте — по этажу
# (`mergeLevelsByFloor`) и по имени операции (`usedOpsForLevel`/`rowsData`,
# app/static/chess-flat.js) — тот же расчёт, что у печати и предпросмотра.
# Здесь ТОЛЬКО отрисовка уже готовых данных в файл, без повторного
# вычисления группировки: вторая (печать) и третья (эта выгрузка)
# реализация одной и той же логики разошлись бы при первой же правке
# любой из них — та же группировка уже дважды была источником живых
# ошибок бланка (Docs/backlog.md 2026-09-15).


def build_walkaround_xlsx(payload: dict) -> bytes:
    """Бланк обхода в Excel — один широкий лист, все секции разом: в
    отличие от печати/PDF, у листа Excel нет физической ширины страницы,
    разбивать секции на группы незачем."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Бланк обхода"

    thin = Side(style="thin", color="9E9E9E")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="F2F2F2")
    absent_fill = PatternFill("solid", fgColor="FAFAFA")

    sections = payload["sections"]
    ncols = 2 + 2 * len(sections)

    ws.append(["Шахматка · %s" % payload["board"]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["%s · уровни %s" % (payload["object_name"], payload["range_label"])])
    ws.append(["Снимок системы: %s" % payload["snapshot_at"], None, "Дата факта:", payload["date"]])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(["Эт.", "Операция"] + [name for name in sections for _ in range(2)])
    ws.append([None, None] + ["В системе", "На дату"] * len(sections))
    ws.merge_cells(start_row=header_row, start_column=1, end_row=header_row + 1, end_column=1)
    ws.merge_cells(start_row=header_row, start_column=2, end_row=header_row + 1, end_column=2)
    for i in range(len(sections)):
        c0 = 3 + i * 2
        ws.merge_cells(start_row=header_row, start_column=c0, end_row=header_row, end_column=c0 + 1)
    for r in (header_row, header_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(bold=True)
            cell.fill = head_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Номер строки ведём сами (не `ws.max_row` в цикле) — тот же приём,
    # что у build_status_report_xlsx: max_row у openpyxl это максимум по
    # всем ячейкам листа, O(n) на каждое обращение.
    row = header_row + 1
    for level_row in payload["rows"]:
        ops = level_row["ops"]
        if not ops:
            row += 1
            ws.cell(row=row, column=1, value=level_row["floor"])
            for c in range(1, ncols + 1):
                ws.cell(row=row, column=c).border = border
            continue
        floor_row = row + 1
        for op in ops:
            row += 1
            ws.cell(row=row, column=2, value=op["name"])
            for i, pct in enumerate(op["cells"]):
                c = 3 + i * 2
                if pct is None:
                    ws.cell(row=row, column=c).fill = absent_fill
                    ws.cell(row=row, column=c + 1).fill = absent_fill
                else:
                    cell = ws.cell(row=row, column=c, value=pct / 100)
                    cell.number_format = "0%"
            for c in range(1, ncols + 1):
                ws.cell(row=row, column=c).border = border
                ws.cell(row=row, column=c).alignment = Alignment(vertical="center")
        ws.cell(row=floor_row, column=1, value=level_row["floor"])
        ws.cell(row=floor_row, column=1).alignment = Alignment(horizontal="center", vertical="center")
        if row > floor_row:
            ws.merge_cells(start_row=floor_row, start_column=1, end_row=row, end_column=1)

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 46
    for i in range(3, ncols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 11
    ws.freeze_panes = ws.cell(row=header_row + 2, column=3)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_walkaround_pdf(payload: dict) -> bytes:
    """Бланк обхода в PDF. По высоте страницы разбивает сам reportlab
    (Platypus меряет свою типографику сам — НЕ пытается повторить
    измерение реального рендера из браузера; те два измерения уже
    расходились ровно на этом бланке при печати, см. Docs/backlog.md
    2026-09-15, «съезжали линии»). По ширине — секции, не помещающиеся на
    одном листе, идут отдельным блоком листов после первых, тот же
    компромисс, что у печати из браузера (`printCapacitySections`,
    app/static/chess-flat.js)."""
    from io import BytesIO

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A3, A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from app.pdf_export import FONT_BOLD, FONT_REGULAR
    from app.reports import pdf_text

    pagesize = A3 if payload.get("format") == "A3" else A4
    page_w, _ = pagesize
    margin = 10 * mm

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=pagesize,
        leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin,
        title="Шахматка · %s" % payload["board"],
    )
    title_style = ParagraphStyle("cf-t", fontName=FONT_BOLD, fontSize=13, leading=16)
    sub_style = ParagraphStyle("cf-s", fontName=FONT_REGULAR, fontSize=9, leading=12,
                               textColor=colors.HexColor("#666666"))
    # Текст ячеек — ОБЯЗАТЕЛЬНО через Paragraph, не голой строкой (живой
    # отчёт пользователя 2026-09-15, скриншот PDF: длинное имя операции
    # вылезало за границу колонки). Table переносит по словам и растягивает
    # высоту строки только у содержимого-Paragraph — голая строка рисуется
    # ОДНОЙ строкой на всю свою естественную длину и просто продолжается
    # поверх соседней колонки, если не влезла (сетка при этом рисуется по
    # заданной colWidths как ни в чём не бывало — несовпадение видно только
    # по тексту).
    op_style = ParagraphStyle("cf-op", fontName=FONT_REGULAR, fontSize=8, leading=10)
    floor_style = ParagraphStyle("cf-floor", fontName=FONT_BOLD, fontSize=9, leading=11,
                                 alignment=TA_CENTER)
    head_style = ParagraphStyle("cf-head", fontName=FONT_BOLD, fontSize=8, leading=10,
                                alignment=TA_CENTER)

    sections = payload["sections"]
    floor_col, op_col, val_col, date_col = 10 * mm, 55 * mm, 14 * mm, 16 * mm
    usable = page_w - 2 * margin
    section_cap = max(1, int((usable - floor_col - op_col) // (val_col + date_col)))

    chunks = []
    i = 0
    if not sections:
        chunks = [([], [])]
    else:
        while i < len(sections):
            chunk = sections[i:i + section_cap]
            chunks.append((chunk, list(range(i, i + len(chunk)))))
            i += len(chunk)

    story = []
    for chunk_i, (chunk, idx) in enumerate(chunks):
        if chunk_i:
            story.append(PageBreak())
        story.append(Paragraph(pdf_text("Шахматка · %s" % payload["board"]), title_style))
        story.append(Paragraph(
            pdf_text("%s · уровни %s" % (payload["object_name"], payload["range_label"])), sub_style))
        story.append(Paragraph(
            pdf_text("Снимок системы: %s     Дата факта: %s" % (payload["snapshot_at"], payload["date"])),
            sub_style))
        story.append(Spacer(1, 4 * mm))

        header1 = ["Эт.", "Операция"]
        for name in chunk:
            header1 += [Paragraph(pdf_text(name), head_style), ""]
        header2 = ["", ""] + ["В системе", "На дату"] * len(chunk)
        data = [header1, header2]
        style_cmds = [("SPAN", (0, 0), (0, 1)), ("SPAN", (1, 0), (1, 1))]
        for i2 in range(len(chunk)):
            c = 2 + i2 * 2
            style_cmds.append(("SPAN", (c, 0), (c + 1, 0)))

        row = 2
        for level_row in payload["rows"]:
            ops = level_row["ops"]
            if not ops:
                data.append([Paragraph(pdf_text(level_row["floor"]), floor_style), ""]
                            + [""] * (2 * len(chunk)))
                row += 1
                continue
            start = row
            for oi, op in enumerate(ops):
                floor_cell = Paragraph(pdf_text(level_row["floor"]), floor_style) if oi == 0 else ""
                line = [floor_cell, Paragraph(pdf_text(op["name"]), op_style)]
                for si in idx:
                    pct = op["cells"][si]
                    line += [("%d%%" % pct) if pct is not None else "", ""]
                data.append(line)
                row += 1
            if row - 1 > start:
                style_cmds.append(("SPAN", (0, start), (0, row - 1)))

        widths = [floor_col, op_col] + [val_col, date_col] * len(chunk)
        table = Table(data, colWidths=widths, repeatRows=2)
        table.setStyle(TableStyle(style_cmds + [
            ("FONTNAME", (0, 0), (-1, -1), FONT_REGULAR),
            ("FONTNAME", (0, 0), (-1, 1), FONT_BOLD),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#F2F2F2")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9E9E9E")),
            ("ALIGN", (2, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(table)

    doc.build(story)
    return buf.getvalue()
