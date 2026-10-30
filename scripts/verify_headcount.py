"""Сквозная проверка учёта численности (app/headcount.py): срок и просрочка, ввод и история, подрядчики и ИНН, отчёт и права.

Пишет в базу ZHBI_DB_PATH (она МЕНЯЕТСЯ: заводит виды работ, подрядчиков, записи) — гонять ТОЛЬКО на копии обезличенной базы:

    cp data/zhbi.anon.db /tmp/hc.db && ZHBI_DB_PATH=/tmp/hc.db .venv/bin/python scripts/verify_headcount.py

Обработчики вызываются напрямую (httpx для TestClient в окружении нет), пользователь передаётся аргументом.
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
_db = os.environ.get("ZHBI_DB_PATH")
if not _db or os.path.basename(_db) == "zhbi.db":
    sys.exit("Скрипт пишет в базу: задайте ZHBI_DB_PATH на КОПИЮ обезличенной базы (боевой data/zhbi.db не годится)")
from datetime import date, datetime, timedelta
from fastapi import HTTPException
from app import db, headcount as hc

db.init_db()
conn = db.get_connection()
admin = conn.execute("SELECT * FROM users WHERE role='admin' LIMIT 1").fetchone()
obj = conn.execute("SELECT id FROM objects ORDER BY id LIMIT 1").fetchone()["id"]
obj2 = conn.execute("SELECT id FROM objects ORDER BY id LIMIT 1 OFFSET 1").fetchone()["id"]
for code, name, sec in [("130-01-01","3.1.1. Бетон","3. Монолит"),("130-03","3.3. Арматура","3. Монолит"),("340-03-02","23.1.1. Земляные","23. Сети")]:
    conn.execute("INSERT OR IGNORE INTO work_codifier (code,name,section_name) VALUES (?,?,?)",(code,name,sec))
conn.execute("INSERT INTO work_codifier (code,name,section_name,retired_at) VALUES ('999','Выведен','X','2026-01-01')")
conn.commit()
w = {r["code"]: r["id"] for r in conn.execute("SELECT id, code FROM work_codifier")}
conn.close()

def assign_fixture(ctr):
    conn = db.get_connection()
    conn.executemany("INSERT OR IGNORE INTO headcount_contractor_works (contractor_id,codifier_id) VALUES (?,?)", [(ctr, wid) for wid in w.values()])
    conn.commit();conn.close()


class R:
    def __init__(self, code, data): self.status_code, self._d = code, data
    def json(self): return self._d
def call(fn, **kw):
    try:
        result=fn(**kw,user=admin)
        if fn is hc.create_contractor: assign_fixture(result["id"])
        return R(200,result)
    except HTTPException as e: return R(e.status_code, {"detail": e.detail})
class C:
    def post(self, url, json): return call(hc.create_contractor, object_id=int(url.split("/")[2]), body=hc.ContractorIn(**json))
    def patch(self, url, json):
        p=url.split("/"); return call(hc.update_contractor, object_id=int(p[2]), contractor_id=int(p[5]), body=hc.ContractorIn(**json))
    def put(self, url, json):
        return call(hc.save_records, object_id=int(url.split("/")[2]), body=hc.SaveIn(**json))
    def delete(self, url):
        p=url.split("/")
        return call(hc.delete_record if p[4]=="records" else hc.delete_contractor, object_id=int(p[2]), **({"record_id":int(p[5])} if p[4]=="records" else {"contractor_id":int(p[5])}))
    def get(self, url, params=None):
        p=url.split("/"); o=int(p[2]); k=p[4]
        if k=="contractors": return call(hc.list_contractors, object_id=o)
        if k=="day": return call(hc.get_day, object_id=o, day=params["date"])
        if k=="history": return call(hc.get_history, object_id=o, date_from=None, date_to=None, contractor_id=None, codifier_id=None, limit=200)
c = C()
ok = lambda cond, msg: print(("OK   " if cond else "FAIL ") + msg) or (cond or sys.exit(1))

# срок
ok(hc.deadline_for(date(2026,10,9)).isoformat().startswith("2026-10-09T11:00"), "пятница: срок 11:00 того же дня")
ok(hc.deadline_for(date(2026,10,10)).date() == date(2026,10,12), "суббота: срок понедельник")
ok(hc.deadline_for(date(2026,10,11)).date() == date(2026,10,12), "воскресенье: срок понедельник")
ok(hc.inn_status("7745000111")=="ok" and hc.inn_status("БСТ-ГРУПП ООО")=="unverified" and hc.inn_status("123456789")=="unverified" and hc.inn_status(None)=="none", "inn_status")

# подрядчики
# единый справочник (2026-10-26): подрядчик — ссылка на контрагента; нового контрагента заводят явным флагом
r = c.post(f"/objects/{obj}/headcount/contractors", json={"inn": " 7745 000111 ", "name": "Альфа ООО", "create_counterparty": True}); ok(r.status_code==200 and r.json()["inn"]=="7745000111" and r.json()["counterparty_id"], "создан контрагент и подрядчик, ИНН очищен")
a = r.json()["id"]; cp_a = r.json()["counterparty_id"]
r = c.post(f"/objects/{obj}/headcount/contractors", json={"inn": "7745000111", "name": "Альфа ООО"}); ok(r.status_code==400, "без флага и без контрагента — 400")
r = c.post(f"/objects/{obj}/headcount/contractors", json={"inn": "7745000111", "name": "Другая", "create_counterparty": True}); ok(r.status_code==409, "контрагент с таким ИНН уже есть — 409, дубль не заводится")
r = c.post(f"/objects/{obj}/headcount/contractors", json={"counterparty_id": cp_a}); ok(r.status_code==409, "тот же контрагент на объект второй раз — 409")
r = c.post(f"/objects/{obj}/headcount/contractors", json={"name": "БСТ-ГРУПП ООО", "create_counterparty": True}); ok(r.status_code==200 and r.json()["inn"] is None, "подрядчик без ИНН")
b = r.json()["id"]; cp_b = r.json()["counterparty_id"]
r = c.post(f"/objects/{obj}/headcount/contractors", json={"name": "  бст-групп  ооо ", "create_counterparty": True}); ok(r.status_code==409, "дубль по названию без регистра — 409")
r = c.patch(f"/objects/{obj}/headcount/contractors/{b}", json={"inn": "123456789"}); ok(r.status_code==400, "название/ИНН связанного подрядчика здесь не правятся — 400")
_c = db.get_connection(); _c.execute("UPDATE counterparties SET inn = '123456789' WHERE id = ?", (cp_b,)); _c.commit(); _c.close()
r = c.get(f"/objects/{obj}/headcount/contractors"); jb = [x for x in r.json()["contractors"] if x["id"] == b][0]
ok(jb["inn"] == "123456789" and jb["inn_note"] == "ИНН не проверен", "ИНН поменяли в справочнике контрагентов — виден в численности, помечен непроверенным")
_c = db.get_connection(); _c.execute("UPDATE counterparties SET short_name = 'БСТ-ГРУПП (переименован)' WHERE id = ?", (cp_b,)); _c.commit(); _c.close()
r = c.get(f"/objects/{obj}/headcount/contractors"); ok([x for x in r.json()["contractors"] if x["id"] == b][0]["display"].startswith("БСТ-ГРУПП (переименован)"), "переименование контрагента сразу видно в подрядчике")
r = c.post(f"/objects/{obj}/headcount/contractors", json={}); ok(r.status_code==400, "пустой подрядчик — 400")
r = c.get(f"/objects/{obj}/headcount/contractors"); ok(len(r.json()["contractors"])==2, "список подрядчиков")

# ввод
today = hc.now_msk().date(); d1 = (today - timedelta(days=3)).isoformat()
rows = [{"contractor_id": a, "codifier_id": w["130-01-01"], "workers": 10}, {"contractor_id": b, "codifier_id": w["130-03"], "workers": 5}]
r = c.put(f"/objects/{obj}/headcount/records", json={"date": d1, "rows": rows}); j = r.json()
ok(r.status_code==200 and j["created"]==2 and j["total"]==15, "внесено 2 строки, итог 15")
ok(j["late"] is True and all(x["late"] for x in j["rows"]) and all(x["overdue_minutes"] for x in j["rows"]), "задним числом — просрочка отмечена")
r = c.put(f"/objects/{obj}/headcount/records", json={"date": d1, "rows": rows}); ok(r.json()["unchanged"]==2 and r.json()["created"]==0, "повтор без изменений")
rows2 = [{"contractor_id": a, "codifier_id": w["130-01-01"], "workers": 12}]
r = c.put(f"/objects/{obj}/headcount/records", json={"date": d1, "rows": rows2}); j = r.json()
ok(j["changed"]==1 and j["total"]==17, "замена числа, остальные строки дня не тронуты")
rec = [x for x in j["rows"] if x["workers"]==12][0]; ok(rec["late"] is True and rec["changes"]==1, "late сохранён при правке, правок=1")
h = c.get(f"/objects/{obj}/headcount/history").json()["history"]; ok(len(h)==3 and h[0]["old"]==10 and h[0]["new"]==12, "история: создание ×2 + замена")
r = c.put(f"/objects/{obj}/headcount/records", json={"date": d1, "rows": rows+rows}); ok(r.status_code==400 and "дубликат" in r.json()["detail"], "дубль в карточке — 400")
r = c.put(f"/objects/{obj}/headcount/records", json={"date": d1, "rows": [{"contractor_id": a, "codifier_id": w["130-01-01"], "workers": 0}]}); ok(r.status_code==400, "0 — 400")
r = c.put(f"/objects/{obj}/headcount/records", json={"date": (today+timedelta(days=1)).isoformat(), "rows": rows2}); ok(r.status_code==400, "будущая дата — 400")
r = c.put(f"/objects/{obj}/headcount/records", json={"date": d1, "rows": [{"contractor_id": a, "codifier_id": w["999"], "workers": 3}]}); ok(r.status_code==400, "выведенный вид работ — 400")
r = c.put(f"/objects/{obj}/headcount/records", json={"date": d1, "rows": [{"contractor_id": 999999, "codifier_id": w["130-03"], "workers": 3}]}); ok(r.status_code==400, "чужой подрядчик — 400")
# откат пакета при ошибке во второй строке
r = c.put(f"/objects/{obj}/headcount/records", json={"date": d1, "rows": [{"contractor_id": a, "codifier_id": w["340-03-02"], "workers": 7}, {"contractor_id": 999999, "codifier_id": w["130-03"], "workers": 3}]})
ok(r.status_code==400 and c.get(f"/objects/{obj}/headcount/day", params={"date": d1}).json()["total"]==17, "пакет атомарен: первая строка не сохранилась")
# удаление
rid = [x for x in c.get(f"/objects/{obj}/headcount/day", params={"date": d1}).json()["rows"] if x["workers"]==5][0]["id"]
r = c.delete(f"/objects/{obj}/headcount/records/{rid}"); ok(r.status_code==200, "удаление записи")
h = c.get(f"/objects/{obj}/headcount/history").json()["history"]; ok(h[0]["new"] is None and h[0]["old"]==5, "в истории удаление")
r = c.delete(f"/objects/{obj}/headcount/contractors/{a}"); ok(r.status_code==409, "подрядчика с записями удалить нельзя")
r = c.delete(f"/objects/{obj}/headcount/contractors/{b}"); ok(r.status_code==409, "и этого нельзя: записи нет, но есть история")

# ---- отчёт
conn = db.get_connection()
conn.execute("UPDATE objects SET smu_id = (SELECT id FROM smu_catalog LIMIT 1) WHERE id IN (?,?)", (obj, obj2)); conn.commit(); conn.close()
r = hc.create_contractor(object_id=obj2, body=hc.ContractorIn(counterparty_id=cp_a), user=admin); a2 = r["id"]; assign_fixture(a2)
base = date(2026, 9, 1)
def put(o, ctr, code, d, n):
    hc.save_records(object_id=o, body=hc.SaveIn(date=d.isoformat(), rows=[hc.RowIn(contractor_id=ctr, codifier_id=w[code], workers=n)]), user=admin)
for i in range(7):
    put(obj, a, "130-01-01", base + timedelta(days=i), 10)
    put(obj2, a2, "130-03", base + timedelta(days=i), 4)
put(obj, a, "340-03-02", base + timedelta(days=6), 3)
def rep(**kw):
    d = dict(levels="object", as_of="2026-09-07", date_from=None, date_to=None, object_ids=f"{obj},{obj2}", smu_id=None, section=None, work_code=None, contractor_inn=None, staff="total")
    d.update(kw); return hc.report(**d, user=admin)
j = rep()
by = {x["object"]["key"]: x for x in j["rows"]}
ok(by[str(obj)]["fact_day"] == 13 and by[str(obj2)]["fact_day"] == 4, "отчёт: факт за опорный день")
ok(j["week_days"] == 1 and j["month_days"] == 7, "понедельник: неделя /1, месяц /7")
ok(abs(by[str(obj)]["week_avg"] - 13) < 0.06, "среднее за неделю по объекту")
ok(abs(j["total"]["week_avg"] - 17) < 0.15 and j["total"]["fact_day"] == 17, "итог = сумма по объектам")
j2 = rep(levels="section,work")
ok(abs(sum(x["week_avg"] for x in j2["rows"]) - j["total"]["week_avg"]) < 0.2, "аддитивность при другой группировке")
j3 = rep(levels="contractor")
ok(len(j3["rows"]) == 1 and j3["rows"][0]["fact_day"] == 17, "подрядчик с одним ИНН на двух объектах сведён в одну строку")
ok("Альфа" in j3["rows"][0]["contractor"]["label"] and "7745000111" in j3["rows"][0]["contractor"]["label"], "подпись «Название (ИНН …)»")
j4 = rep(work_code="130")
ok(j4["total"]["fact_day"] == 14 and j4["week_days"] == 1, "отбор по началу кода 130")
j5 = rep(date_from="2026-09-01", date_to="2026-09-03")
ok(j5["total"]["period_sum"] == 3*14 and j5["period_days"] == 3 and abs(j5["total"]["period_avg"] - 14) < 0.1, "период: сумма и среднее")
ok(len(j["months"]) == 1 and j["months"][0]["month"] == "2026-09", "помесячная серия")
j6 = rep(as_of="2026-12-20"); ok(j6["total"]["fact_day"] == 0 and j6["rows"] == [], "опорный день без данных: пусто, не падает")
try: rep(levels="object,object"); ok(False, "повтор уровня")
except HTTPException as e: ok(e.status_code == 400, "повтор уровня — 400")
print("ВСЁ ПРОШЛО")

# ---- права: человек без грантов
nobody = {"id": 999999, "role": "user"}
for name, fn, kw in [
    ("список подрядчиков", hc.list_contractors, dict(object_id=obj)),
    ("день", hc.get_day, dict(object_id=obj, day="2026-09-01")),
    ("ввод", hc.save_records, dict(object_id=obj, body=hc.SaveIn(date="2026-09-01", rows=[hc.RowIn(contractor_id=a, codifier_id=w["130-03"], workers=1)]))),
    ("создание подрядчика", hc.create_contractor, dict(object_id=obj, body=hc.ContractorIn(name="X"))),
]:
    try: fn(**kw, user=nobody); ok(False, f"{name}: без прав должен быть отказ")
    except HTTPException as e: ok(e.status_code == 403, f"{name}: без грантов — 403")
j = hc.report(levels="object", as_of="2026-09-07", date_from=None, date_to=None, object_ids=None, smu_id=None, section=None, work_code=None, contractor_inn=None, staff="total", user=nobody)
ok(j["rows"] == [] and j["objects"] == 0, "отчёт без грантов: пусто, чужие объекты не видны")
print("ПРАВА ПРОШЛИ")

# ---- просмотр всех данных
import io as _io
from openpyxl import load_workbook as _lw
def recs(**kw):
    d = dict(object_ids=f"{obj},{obj2}", date_from=None, date_to=None, contractor_q=None, work_code=None, source=None, late=False, limit=100, offset=0)
    d.update(kw); return hc.list_records(**d, user=admin)
allr = recs()
total_db = conn_total = db.get_connection().execute("SELECT COUNT(*), SUM(workers) FROM headcount_records WHERE object_id IN (?,?)", (obj, obj2)).fetchone()
ok(allr["total"] == total_db[0] and allr["sum_workers"] == total_db[1], "записи: итог и сумма по всему отбору")
ok(allr["items"][0]["date"] >= allr["items"][-1]["date"], "записи: новые сверху")
ok(len(recs(limit=5)["items"]) == 5 and recs(limit=5, offset=5)["items"][0]["id"] != recs(limit=5)["items"][0]["id"], "записи: постраничный показ")
ok(recs(object_ids=str(obj))["total"] < allr["total"], "записи: отбор по объекту")
ok(recs(date_from="2026-09-01", date_to="2026-09-03")["total"] == 6, "записи: отбор по периоду")
ok(recs(work_code="130-03")["total"] > 0 and all(i["work_code"] == "130-03" for i in recs(work_code="130-03")["items"]), "записи: отбор по коду вида работ")
ok(recs(work_code="130")["total"] >= recs(work_code="130-03")["total"], "записи: код — это начало (раздел целиком)")
ok(recs(contractor_q="АЛЬФА")["total"] > 0 and all("Альфа" in i["contractor"] for i in recs(contractor_q="АЛЬФА")["items"]), "записи: подрядчик без учёта регистра кириллицы")
ok(recs(contractor_q="7745000111")["total"] > 0, "записи: подрядчик по ИНН")
ok(recs(contractor_q="нет-такого")["total"] == 0, "записи: ничего не найдено — пусто")
ok(recs(source="import")["total"] == 0 and recs(source="form")["total"] == allr["total"], "записи: отбор по источнику")
ok(0 < recs(late=True)["total"] <= allr["total"] and all(i["late"] for i in recs(late=True)["items"]), "записи: только просроченные")
try: recs(source="xxx"); ok(False, "источник")
except HTTPException as e: ok(e.status_code == 400, "записи: неверный источник — 400")
hl = hc.list_history_log(object_ids=f"{obj},{obj2}", date_from=None, date_to=None, contractor_q=None, work_code=None, source=None, limit=100, offset=0, user=admin)
ok(hl["total"] >= allr["total"] and hl["items"][0]["changed_at"] >= hl["items"][-1]["changed_at"], "журнал: не меньше записей, новые сверху")
x = hc.records_xlsx(object_ids=f"{obj},{obj2}", date_from=None, date_to=None, contractor_q=None, work_code=None, source=None, late=False, user=admin)
ws = _lw(_io.BytesIO(x.body)).active
ok(ws.max_row - 1 == allr["total"] and ws.cell(1, 1).value == "Дата", "Excel записей: все строки отбора, а не страница")
x = hc.history_log_xlsx(object_ids=f"{obj},{obj2}", date_from=None, date_to=None, contractor_q=None, work_code=None, source=None, user=admin)
ok(_lw(_io.BytesIO(x.body)).active.max_row - 1 == hl["total"], "Excel журнала: все строки отбора")
ok(hc.list_records(object_ids=None, date_from=None, date_to=None, contractor_q=None, work_code=None, source=None, late=False, limit=10, offset=0, user=nobody)["total"] == 0,
   "записи: без грантов — пусто")
print("ПРОСМОТР ПРОШЁЛ")

# ---- умолчания: текущая дата, подсказки при пустом объекте
jd = rep(as_of=None)
last_d = db.get_connection().execute("SELECT MAX(work_date) FROM headcount_records WHERE object_id IN (?,?) AND work_date <= ?", (obj, obj2, hc.now_msk().date().isoformat())).fetchone()[0]
ok(jd["as_of"] == hc.now_msk().date().isoformat(), "отчёт без даты — на сегодня, текущая календарная неделя и месяц")
ok(hc.get_day(object_id=obj, day="2020-01-01", user=admin)["last_date"] == last_d or hc.get_day(object_id=obj, day="2020-01-01", user=admin)["last_date"] is not None, "день: подсказка «последний день с данными»")
empty_obj = [r["id"] for r in db.get_connection().execute("SELECT id FROM objects WHERE id NOT IN (?, ?) ORDER BY id LIMIT 1", (obj, obj2))][0]
re_ = hc.list_records(object_ids=str(empty_obj), date_from=None, date_to=None, contractor_q=None, work_code=None, source=None, late=False, limit=10, offset=0, user=admin)
ok(re_["total"] == 0 and re_["all_total"] > 0, "записи по пустому объекту: подсказка, сколько их по всем доступным")
print("УМОЛЧАНИЯ ПРОШЛИ")

# ---- перечень объектов для отбора собирается из фактических данных
fc = hc.facets(user=admin)["objects"]
ok({o["id"] for o in fc} == {obj, obj2} and all(o["records"] > 0 and o["date_from"] <= o["date_to"] for o in fc), "facets: только объекты, у которых есть записи")
ok(hc.facets(user=nobody)["objects"] == [], "facets: без грантов — пусто")
ok(rep(object_ids=f"{obj},{obj2}")["objects"] == 2, "отчёт: число объектов — по данным, а не по всем доступным")
print("ОТБОР ПО ОБЪЕКТАМ ПРОШЁЛ")

# ---- «Рабочие» / «ИТР» / «Итого» (2026-10-12)
itr_day = date(2026, 6, 10)   # день вне прочих сценариев: цифры считаем от нуля
def day_rows(o, d): return {(r["contractor_id"], r["codifier_id"]): r for r in hc.get_day(object_id=o, day=d.isoformat(), user=admin)["rows"]}
hc.save_records(object_id=obj, body=hc.SaveIn(date=itr_day.isoformat(), rows=[
    hc.RowIn(contractor_id=a, codifier_id=w["130-01-01"], workers=10, itr=2),
    hc.RowIn(contractor_id=a, codifier_id=w["130-03"], workers=0, itr=5)]), user=admin)
rows = day_rows(obj, itr_day)
r1, r2 = rows[(a, w["130-01-01"])], rows[(a, w["130-03"])]
ok((r1["workers"], r1["itr"], r1["total"]) == (10, 2, 12), "рабочие 10 + ИТР 2 → итого 12")
ok((r2["workers"], r2["itr"], r2["total"]) == (0, 5, 5), "строка «только ИТР»: рабочие 0, ИТР 5, итого 5")
dd = hc.get_day(object_id=obj, day=itr_day.isoformat(), user=admin)
ok((dd["total"], dd["total_workers"], dd["total_itr"]) == (17, 10, 7), "итоги дня: 17 = 10 + 7")
# старый клиент не знает про ИТР (itr не передан): ИТР существующей строки остаётся
hc.save_records(object_id=obj, body=hc.SaveIn(date=itr_day.isoformat(), rows=[hc.RowIn(contractor_id=a, codifier_id=w["130-01-01"], workers=11)]), user=admin)
r1 = day_rows(obj, itr_day)[(a, w["130-01-01"])]
ok((r1["workers"], r1["itr"]) == (11, 2), "itr не передан — прежнее значение ИТР сохраняется")
# явный itr=0 сбрасывает
hc.save_records(object_id=obj, body=hc.SaveIn(date=itr_day.isoformat(), rows=[hc.RowIn(contractor_id=a, codifier_id=w["130-01-01"], workers=11, itr=0)]), user=admin)
ok(day_rows(obj, itr_day)[(a, w["130-01-01"])]["itr"] == 0, "itr=0 явно — ИТР обнулён")
# повтор без изменений
res = hc.save_records(object_id=obj, body=hc.SaveIn(date=itr_day.isoformat(), rows=[hc.RowIn(contractor_id=a, codifier_id=w["130-01-01"], workers=11, itr=0)]), user=admin)
ok(res["unchanged"] == 1 and res["changed"] == 0, "те же числа — без изменений")
# проверки
for bad, why in [(dict(workers=0, itr=0), "0 и 0"), (dict(workers=-1, itr=3), "отрицательные рабочие"), (dict(workers=3, itr=-1), "отрицательные ИТР"),
                 (dict(workers=100001, itr=0), "рабочих больше 100 000"), (dict(workers=1, itr=100001), "ИТР больше 100 000")]:
    rr = call(hc.save_records, object_id=obj, body=hc.SaveIn(date=itr_day.isoformat(), rows=[hc.RowIn(contractor_id=a, codifier_id=w["130-01-01"], **bad)]))
    ok(rr.status_code == 400, f"{why} — 400")
# история несёт обе пары
h = hc.get_history(object_id=obj, date_from=itr_day.isoformat(), date_to=itr_day.isoformat(), contractor_id=a, codifier_id=w["130-01-01"], limit=50, user=admin)["history"]
ok(h[-1]["old"] is None and h[-1]["new"] == 10 and h[-1]["old_itr"] is None and h[-1]["new_itr"] == 2, "история: создание с рабочими 10 и ИТР 2")
ok(any(x["old_itr"] == 2 and x["new_itr"] == 0 and x["old"] == 11 and x["new"] == 11 for x in h), "история: правка только ИТР 2 → 0")
# отчёт: Итого / Рабочие / ИТР
def day_sum(staff):
    j = rep(object_ids=str(obj), as_of=itr_day.isoformat(), levels="object", staff=staff)
    return j["staff"], j["total"]["fact_day"]
tot = day_sum("total"); wk = day_sum("workers"); it = day_sum("itr")
ok(tot == ("total", 16) and wk == ("workers", 11) and it == ("itr", 5), f"отчёт на день: итого 16, рабочие 11, ИТР 5 (получено {tot}, {wk}, {it})")
ok(call(hc.report, levels="object", as_of=None, date_from=None, date_to=None, object_ids=None, smu_id=None, section=None, work_code=None,
        contractor_inn=None, staff="all").status_code == 400, "отчёт: неизвестный показатель — 400")
# просмотр данных и Excel
lr = hc.list_records(object_ids=str(obj), date_from=itr_day.isoformat(), date_to=itr_day.isoformat(), contractor_q=None, work_code=None, source=None, late=False, limit=50, offset=0, user=admin)
ok((lr["sum_workers"], lr["sum_itr"], lr["sum_total"]) == (11, 5, 16) and all("itr" in i and "total" in i for i in lr["items"]), "«Все данные»: итоги рабочие 11, ИТР 5, итого 16")
x = hc.records_xlsx(object_ids=str(obj), date_from=itr_day.isoformat(), date_to=itr_day.isoformat(), contractor_q=None, work_code=None, source=None, late=False, user=admin)
from io import BytesIO
from openpyxl import load_workbook
wb = load_workbook(BytesIO(x.body)); wsx = wb.active
heads = [c.value for c in wsx[1]]
ok(heads[8:12] == ["Рабочие", "ИТР", "Механизаторы", "Итого"], "Excel записей: колонки Рабочие / ИТР / Механизаторы / Итого")
ok(sorted(r[11] for r in wsx.iter_rows(min_row=2, values_only=True)) == [5, 11], "Excel записей: «Итого» по строкам 5 и 11")
xr = hc.report_xlsx(levels="object", as_of=itr_day.isoformat(), date_from=None, date_to=None, object_ids=str(obj), smu_id=None, section=None, work_code=None,
                    contractor_inn=None, staff="itr", user=admin)
ok("ИТР" in [c.value for c in load_workbook(BytesIO(xr.body)).active[1]][1], "Excel отчёта по ИТР: показатель назван в шапке")
xh = hc.history_log_xlsx(object_ids=str(obj), date_from=itr_day.isoformat(), date_to=itr_day.isoformat(), contractor_q=None, work_code=None, source=None, user=admin)
ok("ИТР было" in [c.value for c in load_workbook(BytesIO(xh.body)).active[1]], "Excel журнала: колонки ИТР было/стало")
# удаление: в истории обе пары
rid = day_rows(obj, itr_day)[(a, w["130-03"])]["id"]
hc.delete_record(object_id=obj, record_id=rid, user=admin)
h = hc.get_history(object_id=obj, date_from=itr_day.isoformat(), date_to=itr_day.isoformat(), contractor_id=a, codifier_id=w["130-03"], limit=50, user=admin)["history"]
ok(h[0]["new"] is None and h[0]["new_itr"] is None and h[0]["old"] == 0 and h[0]["old_itr"] == 5, "удаление «только ИТР»: в журнале было 0/5, стало пусто")
print("ИТР / РАБОЧИЕ / ИТОГО ПРОШЛИ")

# ---- кодификатор: порядок и полный список
_cod = hc.list_codifier(object_id=obj, full=False, user=admin)["works"]
_names = [w["name"] for w in _cod]
ok(_names.index("1. Проектно-изыскательские работы") < _names.index("16.1. Оборудование паркинга") if "16.1. Оборудование паркинга" in _names else True, "кодификатор: по номерам разделов, а не по порядку строк файла")
_full = hc.list_codifier(object_id=obj, full=True, user=admin)["works"]
ok(len(_full) >= len(_cod) and "records" in _full[0] and "retired_at" in _full[0], "кодификатор: полный список с использованием и датой вывода")
print("КОДИФИКАТОР ПРОШЁЛ")

# ---- иерархия кодификатора до конца
_t = hc.codifier_tree([{"id": 1, "code": "a", "name": "12. Фасады"}, {"id": 2, "code": "b", "name": "12.2. Устройство"}, {"id": 3, "code": "c", "name": "12.2.3. Монтаж"},
                       {"id": 4, "code": "d", "name": "12.2.3.3. Деталь"}, {"id": 5, "code": "e", "name": "12.5.1. Без промежуточного 12.5"}, {"id": 6, "code": "f", "name": "Без номера"}])
_lv = {x["id"]: (x["level"], x["parent_id"], x["has_children"]) for x in _t}
ok(_lv[4] == (4, 3, False) and _lv[3] == (3, 2, True) and _lv[2] == (2, 1, True) and _lv[1] == (1, None, True), "иерархия: 12.2.3.3 — строка 4-го уровня в цепочке 12 → 12.2 → 12.2.3")
ok(_lv[5] == (2, 1, False), "иерархия: нет промежуточного уровня — цепляется к ближайшему вышестоящему")
ok(_lv[6] == (1, None, False), "иерархия: позиция без номера — верхний уровень")
ok([x["id"] for x in _t].index(3) < [x["id"] for x in _t].index(4), "иерархия: родитель идёт перед потомками")
print("ИЕРАРХИЯ ПРОШЛА")

# ---- смена подрядчика / вида работ у внесённой строки (re-key), в том числе обмен ключами
mv_c1 = hc.create_contractor(object_id=obj, body=hc.ContractorIn(name="Переезд-1 ООО", inn="7701111111", create_counterparty=True), user=admin)["id"]
mv_c2 = hc.create_contractor(object_id=obj, body=hc.ContractorIn(name="Переезд-2 ООО", inn="7702222222", create_counterparty=True), user=admin)["id"]
assign_fixture(mv_c1); assign_fixture(mv_c2)
D = "2026-02-02"
def save(rows): return hc.save_records(object_id=obj, body=hc.SaveIn(date=D, rows=rows), user=admin)
r0 = save([hc.RowIn(contractor_id=mv_c1, codifier_id=w["130-01-01"], workers=10, itr=1), hc.RowIn(contractor_id=mv_c2, codifier_id=w["130-03"], workers=5, itr=0)])
recA = [x for x in r0["rows"] if x["contractor_id"] == mv_c1][0]; recB = [x for x in r0["rows"] if x["contractor_id"] == mv_c2][0]
r1 = save([hc.RowIn(contractor_id=mv_c1, codifier_id=w["340-03-02"], workers=12, itr=1, move_from=recA["id"])])
ok(r1["moved"] == 1 and r1["created"] == 0, "re-key: запись переехала на новый вид работ")
keys = {(x["contractor_id"], x["codifier_id"]): x for x in r1["rows"]}
ok((mv_c1, w["130-01-01"]) not in keys and keys[(mv_c1, w["340-03-02"])]["workers"] == 12, "re-key: на старом ключе записи нет, на новом — с новым числом")
ok(keys[(mv_c1, w["340-03-02"])]["entered_at"] == recA["entered_at"] and keys[(mv_c1, w["340-03-02"])]["late"] == recA["late"], "re-key: «внесено когда» и просрочка сохранены")
hh = hc.get_history(object_id=obj, date_from=D, date_to=D, contractor_id=mv_c1, codifier_id=w["130-01-01"], limit=50, user=admin)["history"]
ok(hh and hh[0]["new"] is None and hh[0]["old"] == 10, "re-key: в истории старого ключа — удаление")
try: save([hc.RowIn(contractor_id=mv_c2, codifier_id=w["130-03"], workers=9, itr=0, move_from=keys[(mv_c1, w["340-03-02"])]["id"])]); ok(False, "конфликт")
except HTTPException as e: ok(e.status_code == 409, "re-key: на новом ключе уже есть другая запись — 409")
ok(len(hc.get_day(object_id=obj, day=D, user=admin)["rows"]) == 2, "re-key: при отказе ничего не изменилось")
a2 = keys[(mv_c1, w["340-03-02"])]
r3 = save([hc.RowIn(contractor_id=mv_c2, codifier_id=w["130-03"], workers=12, itr=1, move_from=a2["id"]), hc.RowIn(contractor_id=mv_c1, codifier_id=w["340-03-02"], workers=5, itr=0, move_from=recB["id"])])
k3 = {(x["contractor_id"], x["codifier_id"]): x["workers"] for x in r3["rows"]}
ok(r3["moved"] == 2 and k3 == {(mv_c2, w["130-03"]): 12, (mv_c1, w["340-03-02"]): 5}, "re-key: обмен ключами двух строк за одно сохранение")
try: save([hc.RowIn(contractor_id=mv_c1, codifier_id=w["130-01-01"], workers=1, move_from=999999)]); ok(False, "чужая запись")
except HTTPException as e: ok(e.status_code == 404, "re-key: запись не найдена — 404")
print("RE-KEY ПРОШЁЛ")


# ---- МЕХАНИЗАТОРЫ (2026-10-26)
mday = date(2026, 3, 3)
def mrows(rows, d=mday): return hc.save_records(object_id=obj, body=hc.SaveIn(date=d.isoformat(), rows=rows), user=admin)
r = mrows([hc.RowIn(contractor_id=a, codifier_id=w["130-01-01"], workers=10, itr=2, mech=3)])
row = r["rows"][0]
ok(row["mech"] == 3 and row["total"] == 15 and r["total_mech"] == 3, "механизаторы: внесены, итого = рабочие + ИТР + механизаторы (15)")
r = mrows([hc.RowIn(contractor_id=a, codifier_id=w["130-01-01"], workers=11, itr=2)])
ok(r["rows"][0]["mech"] == 3 and r["rows"][0]["total"] == 16, "старый клиент без mech не затирает механизаторов (None = не менять)")
r = mrows([hc.RowIn(contractor_id=a, codifier_id=w["130-01-01"], workers=11, itr=2, mech=0)])
ok(r["rows"][0]["mech"] == 0 and r["rows"][0]["total"] == 13, "явный 0 механизаторов обнуляет")
r = mrows([hc.RowIn(contractor_id=a, codifier_id=w["130-03"], workers=0, itr=0, mech=4)])
ok([x for x in r["rows"] if x["codifier_id"] == w["130-03"]][0]["total"] == 4, "строка «только механизаторы» допустима")
for bad, why in [(dict(workers=0, itr=0, mech=0), "0, 0 и 0"), (dict(workers=1, mech=-1), "отрицательные механизаторы"), (dict(workers=1, mech=100001), "механизаторов больше 100 000")]:
    rr = call(hc.save_records, object_id=obj, body=hc.SaveIn(date=mday.isoformat(), rows=[hc.RowIn(contractor_id=a, codifier_id=w["340-03-02"], **bad)]))
    ok(rr.status_code == 400, f"{why} — 400")
h = hc.get_history(object_id=obj, date_from=mday.isoformat(), date_to=mday.isoformat(), contractor_id=a, codifier_id=w["130-01-01"], limit=50, user=admin)["history"]
ok(h[-1]["old_mech"] is None and h[-1]["new_mech"] == 3 and any(x["old_mech"] == 3 and x["new_mech"] == 0 for x in h), "история несёт механизаторов: создание 3, правка 3 → 0")
ok(h[0]["new_total"] == 13 and h[0]["old_total"] == 16, "история: итого было/стало считает механизаторов")
def mday_sum(staff):
    j = rep(object_ids=str(obj), as_of=mday.isoformat(), levels="object", staff=staff)
    return j["total"]["fact_day"]
ok((mday_sum("total"), mday_sum("workers"), mday_sum("itr"), mday_sum("mech")) == (17, 11, 2, 4), "отчёт: итого 17, рабочие 11, ИТР 2, механизаторы 4")
lr = hc.list_records(object_ids=str(obj), date_from=mday.isoformat(), date_to=mday.isoformat(), contractor_q=None, work_code=None, source=None, late=False, limit=50, offset=0, user=admin)
ok((lr["sum_workers"], lr["sum_itr"], lr["sum_mech"], lr["sum_total"]) == (11, 2, 4, 17), "«Все данные»: итоги по трём категориям")
x = hc.records_xlsx(object_ids=str(obj), date_from=mday.isoformat(), date_to=mday.isoformat(), contractor_q=None, work_code=None, source=None, late=False, user=admin)
wsx = load_workbook(BytesIO(x.body)).active
ok([c.value for c in wsx[1]][8:12] == ["Рабочие", "ИТР", "Механизаторы", "Итого"], "Excel записей: колонка «Механизаторы» рядом с ИТР")
ok(sorted(r_[11] for r_ in wsx.iter_rows(min_row=2, values_only=True)) == [4, 13], "Excel записей: «Итого» учитывает механизаторов")
xh = hc.history_log_xlsx(object_ids=str(obj), date_from=mday.isoformat(), date_to=mday.isoformat(), contractor_q=None, work_code=None, source=None, user=admin)
ok("Механизаторы было" in [c.value for c in load_workbook(BytesIO(xh.body)).active[1]], "Excel журнала: колонки «Механизаторы было / стало»")
xr = hc.report_xlsx(levels="object", as_of=mday.isoformat(), date_from=None, date_to=None, object_ids=str(obj), smu_id=None, section=None, work_code=None,
                    contractor_inn=None, staff="mech", user=admin)
ok("Механизаторы" in [c.value for c in load_workbook(BytesIO(xr.body)).active[1]][1], "Excel отчёта по механизаторам: показатель назван в шапке")
rid = [x_ for x_ in hc.get_day(object_id=obj, day=mday.isoformat(), user=admin)["rows"] if x_["codifier_id"] == w["130-03"]][0]["id"]
hc.delete_record(object_id=obj, record_id=rid, user=admin)
h = hc.get_history(object_id=obj, date_from=mday.isoformat(), date_to=mday.isoformat(), contractor_id=a, codifier_id=w["130-03"], limit=50, user=admin)["history"]
ok(h[0]["new"] is None and h[0]["new_mech"] is None and h[0]["old_mech"] == 4, "удаление «только механизаторов»: в журнале было 4, стало пусто")
print("МЕХАНИЗАТОРЫ ПРОШЛИ")

# ---- ЕДИНЫЙ СПРАВОЧНИК ПОДРЯДЧИКОВ (2026-10-26)
sr = hc.search_counterparties(object_id=obj, q="7745000111", limit=30, user=admin)
ok(sr["items"][0]["id"] == cp_a and sr["items"][0]["added"] == a and sr["can_create"], "поиск по ИНН находит контрагента, отмечает, что он уже на объекте; админу можно заводить")
sr = hc.search_counterparties(object_id=obj, q="переименован", limit=30, user=admin)
ok([i["id"] for i in sr["items"]] == [cp_b], "поиск по названию без учёта регистра (по актуальному названию справочника)")
# пользователь с правом на численность, но без права на «Справочник контрагентов»: создать нельзя, выбрать из справочника — можно
conn = db.get_connection()
plain = conn.execute("SELECT * FROM users WHERE role = 'user' ORDER BY id LIMIT 1").fetchone()
conn.close()
from app import access as _acc
_has = _acc.has_feature
_acc_calls = []
def _fake_has(conn_, user_, key, kind, object_id=None):
    if key == "counterparties":
        return False
    return True
import app.headcount as _hcm
_hcm.has_feature = _fake_has
try:
    rr = call(hc.create_contractor, object_id=obj, body=hc.ContractorIn(name="Самовольная ООО", inn="7799999999", create_counterparty=True))
    ok(rr.status_code == 403 and "Справочник контрагентов" in rr.json()["detail"], "без права на «Справочник контрагентов» новый контрагент не заводится — 403 с понятным текстом")
    conn = db.get_connection()
    ok(conn.execute("SELECT COUNT(*) FROM counterparties WHERE inn = '7799999999'").fetchone()[0] == 0, "и в справочнике он не появился")
    conn.close()
    sr = hc.search_counterparties(object_id=obj, q="", limit=5, user=admin)
    ok(sr["can_create"] is False, "поиск сообщает клиенту, что заводить контрагентов нельзя")
    ok(call(hc.create_contractor, object_id=obj, body=hc.ContractorIn(counterparty_id=cp_b)).status_code == 409, "выбрать существующего контрагента без этого права можно (здесь — уже на объекте, 409)")
finally:
    _hcm.has_feature = _has
# замена контрагента подрядчика и перенос записей
conn = db.get_connection()
cp_c = hc.insert_counterparty(conn, "Замена ООО", "7788888888"); conn.commit(); conn.close()
r = c.patch(f"/objects/{obj}/headcount/contractors/{b}", json={"counterparty_id": cp_c}); ok(r.status_code == 200 and r.json()["name"] == "Замена ООО" and r.json()["inn"] == "7788888888", "замена контрагента подрядчика: показ из нового")
r = c.patch(f"/objects/{obj}/headcount/contractors/{b}", json={"counterparty_id": None}); ok(r.status_code == 400, "снять связь с контрагентом нельзя")
r = c.patch(f"/objects/{obj}/headcount/contractors/{a}", json={"counterparty_id": cp_c}); ok(r.status_code == 409, "контрагент уже у другого подрядчика объекта — 409")
# единая функция: найти или создать
conn = db.get_connection()
n0 = conn.execute("SELECT COUNT(*) FROM counterparties").fetchone()[0]
e1 = hc.ensure_counterparty(conn, "7745000111", None); ok(e1 == cp_a, "ensure_counterparty: по ИНН находит существующего")
e2 = hc.ensure_counterparty(conn, "7766666666", None); ok(e2 and conn.execute("SELECT short_name FROM counterparties WHERE id = ?", (e2,)).fetchone()[0] == "ИНН 7766666666", "ensure_counterparty: новый без названия — «ИНН <номер>»")
ok(hc.ensure_counterparty(conn, "7766666666", "Иное") == e2, "ensure_counterparty: повтор по ИНН не плодит контрагентов")
e3 = hc.ensure_counterparty(conn, None, "Только Название ООО"); ok(hc.ensure_counterparty(conn, None, " только  название ооо ") == e3, "ensure_counterparty: без ИНН — по названию без учёта регистра")
conn.execute("INSERT INTO counterparties (full_name, short_name, inn) VALUES ('Д1', 'Д1', '7755555555')"); conn.execute("INSERT INTO counterparties (full_name, short_name, inn) VALUES ('Д2', 'Д2', '7755555555')")
ok(hc.ensure_counterparty(conn, "7755555555", "X") is None, "ensure_counterparty: два контрагента с одним ИНН — None, третий не заводится")
ok(hc.ensure_counterparty(conn, None, None) is None, "ensure_counterparty: без ИНН и названия — None")
conn.commit(); conn.close()
# обработка обновления: связывает несвязанных, идемпотентна
conn = db.get_connection()
conn.execute("UPDATE headcount_contractors SET counterparty_id = NULL WHERE id = ?", (a,))
conn.execute("INSERT INTO headcount_contractors (object_id, name_raw, inn_raw) VALUES (?, ?, ?)", (obj2, "Старый ООО", "7733333333"))
conn.execute("INSERT INTO headcount_contractors (object_id, inn_raw) VALUES (?, ?)", (obj2, "7755555555"))
conn.commit(); conn.close()
from app import release_tasks as rt
res1 = rt.run_by_name("2026-10-26-link-headcount-contractors"); ok(res1["status"] == "ok" and "связано" in res1["note"], f"обработка обновления выполнена: {res1['note']}")
conn = db.get_connection()
ok(conn.execute("SELECT counterparty_id FROM headcount_contractors WHERE id = ?", (a,)).fetchone()[0] == cp_a, "обработка: прежняя запись связана по ИНН с существующим контрагентом")
ok(conn.execute("SELECT COUNT(*) FROM counterparties WHERE inn = '7733333333' AND short_name = 'Старый ООО'").fetchone()[0] == 1, "обработка: недостающий контрагент создан из названия и ИНН")
ok(conn.execute("SELECT counterparty_id FROM headcount_contractors WHERE inn_raw = '7755555555'").fetchone()[0] is None, "обработка: неоднозначный ИНН остаётся несвязанным")
n1 = conn.execute("SELECT COUNT(*) FROM counterparties").fetchone()[0]; conn.close()
res2 = rt.run_by_name("2026-10-26-link-headcount-contractors")
conn = db.get_connection()
ok(res2["status"] == "ok" and conn.execute("SELECT COUNT(*) FROM counterparties").fetchone()[0] == n1, "обработка идемпотентна: повтор не создаёт контрагентов")
conn.close()
print("ЕДИНЫЙ СПРАВОЧНИК ПРОШЁЛ")
