"""Смены, календарные средние, независимые сводки и план монтажа. Только тестовая копия БД."""
import os
import sys
from pathlib import Path
from datetime import date, datetime
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
if not os.environ.get("ZHBI_DB_PATH") or Path(os.environ["ZHBI_DB_PATH"]).name == "zhbi.db":
    sys.exit("Нужна тестовая копия в ZHBI_DB_PATH")
from fastapi import HTTPException
from app import db, headcount as hc, express_bot as eb, express_summaries as summaries, express_scheduler as scheduler
from app import release_tasks
from app.element_dates import set_planned_installation_dates
from app.db_schema_doc import verify


def check(value, text):
    assert value, text
    print("OK " + text)


c=db.get_connection()
before=[tuple(r) for r in c.execute("SELECT id,object_id,work_date,contractor_id,codifier_id,workers,itr,mech,entered_at FROM headcount_records ORDER BY id")]
c.close()
db.init_db()
c=db.get_connection()
after=[tuple(r) for r in c.execute("SELECT id,object_id,work_date,contractor_id,codifier_id,workers,itr,mech,entered_at FROM headcount_records ORDER BY id")]
check(before==after and not c.execute("SELECT 1 FROM headcount_records WHERE shift<>'day'").fetchone(),"миграция сохраняет данные, id и даты; прежние записи дневные")
check(not [x for x in verify(c) if any(t in x for t in ('express_','headcount_','planned_installation'))],"новые поля, таблицы и связи описаны")
admin=c.execute("SELECT * FROM users WHERE role='admin' LIMIT 1").fetchone()
obj=c.execute("SELECT id FROM objects ORDER BY id LIMIT 1").fetchone()[0]
ctr=c.execute("INSERT INTO headcount_contractors (object_id,name_raw,inn_raw) VALUES (?,'Сменный подрядчик','9999999988')",(obj,)).lastrowid
wid=c.execute("INSERT INTO work_codifier (code,name,section_name) VALUES ('SHIFT-1','880003. Сменная работа','Сменный раздел')").lastrowid
c.execute("INSERT INTO headcount_contractor_works (contractor_id,codifier_id,note) VALUES (?,?,'Каркас')",(ctr,wid))
c.commit();c.close()
fixed=datetime(2026,10,30,13,tzinfo=hc.MSK)
with patch.object(hc,'now_msk',return_value=fixed):
    for shift,values in [('day',(10,2,1)),('night',(20,3,2))]:
        hc.save_records(obj,hc.SaveIn(date='2026-10-30',shift=shift,rows=[hc.RowIn(contractor_id=ctr,codifier_id=wid,workers=values[0],itr=values[1],mech=values[2])]),user=admin)
    hc.save_row_from_bot(obj,date(2026,10,30),ctr,wid,25,4,3,contact_name='Тест',shift='night')
    day=next(x for x in hc.get_day(obj,'2026-10-30',user=admin,shift='day')['rows'] if x['contractor_id']==ctr)
    night=next(x for x in hc.get_day(obj,'2026-10-30',user=admin,shift='night')['rows'] if x['contractor_id']==ctr)
    check(day['total']==13 and night['total']==32,'форма и бот сохраняют смены отдельно и не перезаписывают другую смену')
    hc.save_row_from_bot(obj,date(2026,10,26),ctr,wid,7,0,0,contact_name='Тест')
    hc.save_row_from_bot(obj,date(2026,10,1),ctr,wid,8,0,0,contact_name='Тест')
    args=dict(levels='contractor,section,note',as_of='2026-10-30',date_from='2026-10-26',date_to='2026-10-30',object_ids=str(obj),smu_id=None,section=None,work_code=None,contractor_inn='9999999988',staff='total',user=admin)
    report=hc.report(**args)
    check(report['week_from']=='2026-10-26' and report['week_days']==5 and report['total']['week_avg']==10.4,'пятница: текущая неделя /5, включая дни без работы')
    check(report['month_days']==30 and report['total']['month_avg']==2,'октябрь по 30 число: обе смены /30')
    check(report['total']['period_sum']==52 and report['total']['period_avg']==10.4,'произвольный период: все календарные дни включительно')
    check([x['total'] for x in report['series']]==[7,0,0,0,45],'диаграмма содержит все дни и обе смены')
    c=db.get_connection()
    check([tuple(r) for r in c.execute("SELECT shift,old_workers,new_workers FROM headcount_history WHERE record_id=? ORDER BY id",(night['id'],))]==[('night',None,20),('night',20,25)],'история ночной смены изолирована')
    contacts=[]
    for i in (1,2):
        contacts.append(c.execute("INSERT INTO express_contacts (name,chat_id) VALUES (?,?)",(f'Получатель {i}',f'summary-chat-{i}')).lastrowid)
    legacy=c.execute("INSERT INTO express_schedules (object_id,contact_id,weekdays,send_time,summary_time,summary_contact_id,effective_from) VALUES (?,?,'1,2,3,4,5','10:30','12:30',?,'2000-01-01 00:00:00')",(obj,contacts[0],contacts[0])).lastrowid
    c.execute("INSERT INTO express_requests (schedule_id,object_id,contact_id,work_date,summary_at) VALUES (?,?,?,'2026-10-29','2026-10-29 09:30:00')",(legacy,obj,contacts[0]))
    c.commit()
    release_tasks._split_express_summaries(c);c.commit()
    check(c.execute("SELECT summary_time FROM express_schedules WHERE id=?",(legacy,)).fetchone()[0] is None,'прежняя сводка перенесена отдельно от опроса')
    check(c.execute("SELECT status FROM express_summary_deliveries WHERE run_date='2026-10-29' AND contact_id=?",(contacts[0],)).fetchone()[0]=='sent','перенос сводок сохраняет отметки прежних отправок')
    count=c.execute('SELECT COUNT(*) FROM express_summary_schedules').fetchone()[0]
    release_tasks._split_express_summaries(c);c.commit()
    check(count==c.execute('SELECT COUNT(*) FROM express_summary_schedules').fetchone()[0],'перенос сводок идемпотентен')
    c.close()
    schedule=summaries.create_summary_schedule(obj,summaries.SummaryScheduleIn(contact_ids=contacts,send_time='13:00'),user=admin)
    c=db.get_connection();c.execute("UPDATE express_summary_schedules SET effective_from='2000-01-01 00:00:00'");c.commit();c.close()
    sent=[]
    with patch.object(eb,'is_active',return_value=True),patch.object(eb,'send_message',side_effect=lambda *a:sent.append(a)):
        first=scheduler.tick(fixed)
        check({x[0] for x in sent}=={'summary-chat-1','summary-chat-2'},'сводка всего объекта отправлена нескольким получателям без опроса')
        check(any('Сменный подрядчик / Каркас / 35 / 6 / 4' in x[1] for x in sent),'формат сводки: подрядчик / примечание / рабочие / ИТР / механизаторы, обе смены')
        sent.clear();scheduler.tick(fixed)
        check(not sent,'повторный обход не повторяет сводки')
    c=db.get_connection();ids=[r[0] for r in c.execute('SELECT id FROM elements LIMIT 2')]
    c.execute("UPDATE elements SET current_status='ready_for_installation' WHERE id=?",(ids[0],));c.execute("UPDATE elements SET current_status='delivered' WHERE id=?",(ids[1],));c.commit()
    try:
        set_planned_installation_dates(c,ids,'2026-11-01',admin)
        raise AssertionError('смешанные статусы')
    except HTTPException as e:check(e.status_code==409,'план монтажа: смешанная пачка отклоняется целиком')
    check(c.execute('SELECT planned_installation_date FROM elements WHERE id=?',(ids[0],)).fetchone()[0] is None,'при отказе дата готового изделия не меняется')
    set_planned_installation_dates(c,[ids[0]],'2026-11-01',admin);c.commit()
    check(c.execute('SELECT planned_installation_date FROM elements WHERE id=?',(ids[0],)).fetchone()[0]=='2026-11-01','план монтажа сохраняется у готового изделия')
    try:
        set_planned_installation_dates(c,[ids[0]],'2026-02-30',admin)
        raise AssertionError('неверная дата')
    except HTTPException as e:check(e.status_code==400,'неверная дата монтажа отклоняется')
    c.close()
check(not db.init_db(),'миграции повторно ничего не меняют')

# Настоящие маршруты и ограничения доступа, без запуска приложения/планировщика.
from fastapi import FastAPI
import asyncio
import json
from app.auth import get_current_user
from app.main import update_installation_date
api = FastAPI()
api.include_router(summaries.router)
api.add_api_route('/elements/bulk-planned-installation-date', update_installation_date, methods=['PATCH'])
viewer = dict(admin); viewer.update(id=-12345, role='viewer', is_system_admin=0)
current_user = [viewer]
api.dependency_overrides[get_current_user] = lambda: current_user[0]
def request(method, path, payload):
    messages = []
    async def run():
        async def receive():
            return {'type':'http.request','body':json.dumps(payload).encode(),'more_body':False}
        async def send(message):
            messages.append(message)
        await api({'type':'http','asgi':{'version':'3.0'},'http_version':'1.1','method':method,
                   'scheme':'http','path':path,'raw_path':path.encode(),'query_string':b'',
                   'headers':[(b'content-type',b'application/json')], 'server':('test',80),'client':('test',1)}, receive, send)
    asyncio.run(run())
    return next(m['status'] for m in messages if m['type']=='http.response.start')
check(request('PATCH','/elements/bulk-planned-installation-date', {'element_ids':[ids[0]], 'planned_installation_date':'2026-11-02'}) == 403,
      'HTTP: без права изменения плановая дата не сохраняется')
check(request('POST',f'/objects/{obj}/express/summary-schedules', {'contact_ids':contacts,'send_time':'14:00'}) == 403,
      'HTTP: без права записи расписание сводки не создаётся')
current_user[0] = admin
check(request('PATCH','/elements/bulk-planned-installation-date', {'element_ids':[ids[0]], 'planned_installation_date':None}) == 200,
      'HTTP: плановую дату готового изделия можно снять')
check(request('PATCH','/elements/bulk-planned-installation-date', {'element_ids':[ids[1]], 'planned_installation_date':'2026-11-02'}) == 409,
      'HTTP: дата запрещена вне статуса готовности к монтажу')
check(request('POST',f'/objects/{obj}/express/summary-schedules', {'contact_ids':[],'send_time':'14:00'}) == 422,
      'HTTP: пустой список получателей отклоняется')

# Смена различима и в выгрузке: проверяем содержимое, затем структуру XLSX отдельным инструментом.
from io import BytesIO
from openpyxl import load_workbook
export_args = dict(object_ids=str(obj),date_from='2026-10-30',date_to='2026-10-30',contractor_q='9999999988',work_code=None,source=None,user=admin)
for name, response in [('records',hc.records_xlsx(**export_args,late=False)),('history',hc.history_log_xlsx(**export_args))]:
    wb = load_workbook(BytesIO(response.body)); ws = wb.active
    shift_col = [c.value for c in ws[1]].index('Смена')
    check({r[shift_col] for r in ws.iter_rows(min_row=2,values_only=True)} == {'Дневная','Ночная'}, f'Excel {name}: обе смены названы явно')
    output = Path('/tmp') / f'headcount-shifts-{name}.xlsx'
    output.write_bytes(response.body)
    import subprocess
    subprocess.run([sys.executable,'/Users/max/.codex/scripts/validate_xlsx.py',str(output)],check=True)
print('Все проверки пройдены')
